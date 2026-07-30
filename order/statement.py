import io, datetime
import logging
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from .models import ClothOrder

FONT   = Font(name="\u7b49\u7ebf", size=11)
FBOLD  = Font(name="\u7b49\u7ebf", size=11, bold=True)
FTITLE = Font(name="\u7b49\u7ebf", size=16)
FSUB   = Font(name="\u7b49\u7ebf", size=14)
FQTY   = Font(name="\u5b8b\u4f53", size=11, bold=True)
AC     = Alignment(horizontal="center", vertical="center", wrap_text=True)
BDR    = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

HD = ["\u8ba2\u5355\u5e8f\u53f7", "\u9001\u8d27\u5355\u53f7", "\u54c1\u540d", "\u989c\u8272",
      "\u51fa\u8d27\u65e5\u671f", "\u6570\u91cf\uff08\u7c73\uff09", "\u5355\u4ef7", "\u5355\u4f4d",
      "\u91d1\u989d", "\u62a5\u4ef7\u5f62\u5f0f", "\u7ec6\u6570"]
WD = [8.875, 15.875, 11.625, 10.216, 10.875, 10.625, 10.875, 11.0, 13.0, 20.625, 13.0]
NC = len(HD)


def _setup_sheet(ws):
    for i, w in enumerate(WD, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, h in [(1,20.25),(2,18),(4,18),(5,21)]:
        ws.row_dimensions[r].height = h
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NC)
    c = ws.cell(row=1, column=1, value="\u5e7f\u5dde\u5e02\u6607\u5955\u9686\u7eba\u7ec7\u54c1\u6709\u9650\u516c\u53f8")
    c.font = FTITLE; c.alignment = AC
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NC)
    c = ws.cell(row=2, column=1, value="\u5bf9\u8d26\u5355")
    c.font = FSUB; c.alignment = AC
    ws.cell(row=4, column=2, value="BTC").font = FBOLD; ws.cell(row=4, column=2).alignment = AC
    today = datetime.date.today()
    ws.merge_cells(start_row=4, start_column=9, end_row=4, end_column=10)
    c = ws.cell(row=4, column=9, value="\u5bf9\u8d26\u65e5\u671f\uff1a" + today.strftime("%Y-%m-%d"))
    c.font = FONT; c.alignment = Alignment(horizontal="right", vertical="center")
    for col, h in enumerate(HD, 1):
        c = ws.cell(row=5, column=col, value=h)
        c.font = FONT; c.alignment = AC; c.border = BDR

def _add_order_rows(ws, order, start_row, date_seq=None):
    row = start_row
    tq = 0
    if date_seq is None:
        date_seq = {}
    for s in order.shipments.filter(is_deleted=False).order_by("batch_number"):
        dt = s.date
        qv = s.quantity
        if not dt and not qv: continue
        qv = float(qv) if qv else 0
        pr = float(order.price) if order.price else 0
        tq += qv
        if dt:
            date_key = dt.isoformat()
            date_seq[date_key] = date_seq.get(date_key, 0) + 1
            seq = date_seq[date_key]
        else:
            seq = 1
        dn = f"SYL{dt.strftime('%Y%m%d')}-{seq}" if dt else ""
        ws.cell(row=row, column=1,  value=order.serial_number or "").font = FONT
        ws.cell(row=row, column=2,  value=dn).font = FONT
        ws.cell(row=row, column=3,  value=order.cloth_type or "").font = FONT
        ws.cell(row=row, column=4,  value=order.color or "").font = FONT
        ws.cell(row=row, column=5,  value=f"{dt.month}\u6708{dt.day}\u65e5" if dt else "").font = FONT
        ws.cell(row=row, column=6,  value=qv).font = FQTY
        ws.cell(row=row, column=7,  value=pr).font = FQTY
        unit = order.price_unit or "元/米"
        ws.cell(row=row, column=8,  value=unit).font = FONT
        ws.cell(row=row, column=9).value = round(pr * qv, 2)
        ws.cell(row=row, column=9).font = FONT
        ws.cell(row=row, column=9).number_format = '#,##0.00'
        ws.cell(row=row, column=10).font = FONT
        ws.cell(row=row, column=11).font = FONT
        for c in range(1, NC+1):
            ws.cell(row=row, column=c).alignment = AC
            ws.cell(row=row, column=c).border = BDR
        ws.row_dimensions[row].height = 21
        row += 1
    return row, tq

def _add_image(ws, row):
    try:
        from openpyxl.drawing.image import Image
        from django.conf import settings
        img_path = settings.STATIC_ROOT / "img/invoice-details.png"
        if img_path.exists():
            img = Image(str(img_path))
            img.width = 700
            img.height = int(img.width * 0.6)
            ws.add_image(img, f"C{row + 3}")
        else:
            logging.info("Unable to find the invoicing image file, skipping generation of image")
    except Exception as e:
        logging.error(f" {e}", exc_info=True)

def generate_order_statement(order):
    return generate_bulk_statement([order])

def generate_bulk_statement(orders):
    wb = Workbook()
    ws = wb.active
    ws.title = "对账单"
    _setup_sheet(ws)
    date_seq = {}
    row = 6
    dr_start = row
    grand_qty = 0
    total_money = 0.0
    
    for order in orders:
        row, oq = _add_order_rows(ws, order, row, date_seq)
        grand_qty += oq

    if grand_qty > 0:
        for r in range(dr_start, row):
            val = ws.cell(row=r, column=9).value
            if isinstance(val, (int, float)):
                total_money += val
                
    if not grand_qty:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NC)
        ws.cell(row=row, column=1, value="暂无出货记录").font = FONT
        ws.cell(row=row, column=1).alignment = AC
        row += 1
    else:
        row += 1

    ws.cell(row=row, column=5, value="合计").font = FBOLD
    ws.cell(row=row, column=6, value=grand_qty).font = FBOLD
    ws.cell(row=row, column=9, value=round(total_money, 2)).font = FBOLD
    
    ws.cell(row=row, column=9).number_format = '¥#,##0.00'
    
    for c in range(1, NC+1):
        ws.cell(row=row, column=c).alignment = AC
        ws.cell(row=row, column=c).border = BDR
        
    _add_image(ws, row)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

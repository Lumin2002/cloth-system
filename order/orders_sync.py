"""订单与本地 Excel 文件的双向同步。

设计约定：
1. Django 数据库仍是主数据源；本地 Excel 是「工作副本」。
2. Django 侧编辑后，通过信号防抖地重新生成同步 Excel。
3. 后台 watcher 轮询同步 Excel 的 hash；发现外部修改时导入数据库。
4. 已取消订单在 Excel 中通过整行删除线表示，和现有导入逻辑保持一致。

局限性（使用前请知悉）：
- 请尽量单人、单机编辑该 Excel，避免两个人同时改造成冲突。
- Excel/WPS 打开文件时会加锁，保存并关闭后同步才会生效。
- 首次发现已有 Excel 文件时不会自动覆盖数据库；可用
  `python manage.py sync_orders --import-now` 主动导入一次。
"""

import hashlib
import json
import logging
import os
import socket
import threading
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
from django.conf import settings

from .order_excel import EXCEL_COLUMN_MAP, _display_value, import_orders_from_file
from .models import ClothOrder, Shipment

logger = logging.getLogger(__name__)

# 排除掉重复字段的别名列，导出只保留规范表头。
SYNC_COLUMNS = [
    (header, field)
    for header, field in EXCEL_COLUMN_MAP.items()
    if header not in {"*订单类似", "是否开票.1"}
]

_SYNC_STATE_FILENAME = "orders_sync_state.json"
_export_timer = None
_export_lock = threading.Lock()
_state_lock = threading.Lock()
_sync_context = threading.local()


def get_sync_enabled() -> bool:
    return bool(getattr(settings, "ORDER_EXCEL_SYNC_ENABLED", False))


def get_sync_file_path() -> Path:
    path = getattr(settings, "ORDER_EXCEL_SYNC_FILE", None)
    if not path:
        path = Path(settings.BASE_DIR) / "sync" / "orders_sync.xlsx"
    return Path(path).resolve()


def get_sync_interval() -> float:
    value = getattr(settings, "ORDER_EXCEL_SYNC_INTERVAL", 3)
    try:
        return max(float(value), 1.0)
    except (TypeError, ValueError):
        return 3.0


def _state_path() -> Path:
    return get_sync_file_path().with_name(_SYNC_STATE_FILENAME)


def _read_state() -> dict:
    path = _state_path()
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        logger.exception("[Excel同步] 读取同步状态失败，按空状态处理")
        return {}


def _write_state(state: dict) -> None:
    path = _state_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, path)


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def build_sync_dataframe(queryset=None) -> pd.DataFrame:
    """生成与导入格式完全对应的全量同步 DataFrame。"""
    if queryset is None:
        queryset = ClothOrder.objects.all().order_by("serial_number")
    if hasattr(queryset, "prefetch_related"):
        queryset = queryset.prefetch_related("shipments")

    shipments = (
        Shipment.objects.filter(order__in=queryset, is_deleted=False)
        .order_by("order_id", "batch_number")
    )
    shipments_by_order = {}
    for shipment in shipments:
        shipments_by_order.setdefault(shipment.order_id, []).append(shipment)

    rows = []
    for order in queryset:
        values = {
            field: _display_value(order, field)
            for _, field in SYNC_COLUMNS
        }
        order_shipments = shipments_by_order.get(order.pk, [])

        for index in range(1, 6):
            shipment = order_shipments[index - 1] if index <= len(order_shipments) else None
            values[f"finished_product_shipment_date_{index}"] = (
                shipment.date if shipment else None
            )
            values[f"finished_product_shipment_quantity_{index}"] = (
                float(shipment.quantity)
                if shipment and shipment.quantity is not None
                else None
            )

        rows.append(values)

    df = pd.DataFrame(rows)
    if df.empty:
        df = pd.DataFrame(columns=[field for _, field in SYNC_COLUMNS])

    df = df[[field for _, field in SYNC_COLUMNS]]
    df.columns = [header for header, _ in SYNC_COLUMNS]
    return df


def _cancelled_flags(queryset) -> list:
    return [order.order_status == "cancelled" for order in queryset]


def _apply_strikethrough(path: Path, cancelled_flags: list) -> None:
    if not cancelled_flags:
        return
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    wb = load_workbook(path)
    ws = wb["大货"] if "大货" in wb.sheetnames else wb.worksheets[0]
    for row_index, cancelled in enumerate(cancelled_flags, start=2):
        if not cancelled:
            continue
        for cell in ws[row_index]:
            original = cell.font
            cell.font = Font(
                name=original.name,
                size=original.size,
                bold=original.bold,
                italic=original.italic,
                underline=original.underline,
                strike=True,
                color=original.color,
            )
    wb.save(path)


def export_sync_file(queryset=None) -> Path:
    """将数据库订单全量写入同步 Excel 文件。"""
    if queryset is None:
        queryset = ClothOrder.objects.all().order_by("serial_number")
    queryset = list(queryset)

    path = get_sync_file_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.stem + ".sync.tmp.xlsx")

    df = build_sync_dataframe(queryset)
    with pd.ExcelWriter(tmp, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="大货", index=False)

    _apply_strikethrough(tmp, _cancelled_flags(queryset))
    os.replace(tmp, path)

    with _state_lock:
        state = _read_state()
        state["known_hash"] = _sha256_file(path)
        state["last_export_time"] = datetime.now().isoformat(timespec="seconds")
        state["pending_export"] = False
        _write_state(state)

    logger.info(f"[Excel同步] 已从数据库导出 {len(queryset)} 条订单到 {path.name}")
    return path


def _is_importing() -> bool:
    return bool(getattr(_sync_context, "importing", False))


def import_sync_file_if_changed(force: bool = False):
    """检测同步 Excel 是否有外部修改；有变化则导入数据库。"""
    path = get_sync_file_path()
    if not path.exists():
        return None

    with _state_lock:
        state = _read_state()
        pending_export = bool(state.get("pending_export"))
    if pending_export:
        try:
            return export_sync_file()
        except (PermissionError, OSError) as exc:
            _set_pending_export()
            logger.warning(f"[Excel同步] 待导出重试失败：{exc}")
            return None

    try:
        current_hash = _sha256_file(path)
    except Exception as exc:
        logger.warning(f"[Excel同步] 读取文件失败（可能正被 Excel 占用）：{exc}")
        return None

    if not force and current_hash == state.get("known_hash"):
        return None

    _sync_context.importing = True
    try:
        with path.open("rb") as f:
            result = import_orders_from_file(f)
        with _state_lock:
            state = _read_state()
            state["known_hash"] = _sha256_file(path)
            state["last_import_time"] = datetime.now().isoformat(timespec="seconds")
            state["pending_export"] = False
            _write_state(state)
        logger.info(
            f"[Excel同步] 已导入外部修改：新增 {result.imported} 条，"
            f"更新 {result.updated} 条，失败 {result.errors} 条"
        )
        return result
    except Exception as exc:
        logger.exception(f"[Excel同步] 导入外部修改失败：{exc}")
        return None
    finally:
        _sync_context.importing = False


def schedule_export(delay: float = 2.0) -> None:
    """Django 编辑后防抖调度一次导出，避免每个字段保存都写 Excel。"""
    if not get_sync_enabled() or _is_importing():
        return
    global _export_timer
    with _export_lock:
        if _export_timer is not None:
            _export_timer.cancel()
        _export_timer = threading.Timer(delay, _do_export)
        _export_timer.daemon = True
        _export_timer.start()


def _do_export() -> None:
    try:
        export_sync_file()
    except (PermissionError, OSError) as exc:
        _set_pending_export()
        logger.warning(f"[Excel同步] 导出被跳过（文件可能被 Excel 占用），稍后重试：{exc}")
    except Exception:
        logger.exception("[Excel同步] 导出同步 Excel 失败")


def _set_pending_export() -> None:
    with _state_lock:
        state = _read_state()
        state["pending_export"] = True
        _write_state(state)


def _try_bind_lock(port: int):
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        sock.bind(("127.0.0.1", port))
        sock.listen(1)
        return sock
    except OSError:
        try:
            sock.close()
        except Exception:
            pass
        return None


def _ensure_initialized() -> None:
    path = get_sync_file_path()
    with _state_lock:
        state = _read_state()
    if state.get("known_hash"):
        return
    if path.exists():
        try:
            file_hash = _sha256_file(path)
            with _state_lock:
                state = _read_state()
                state["known_hash"] = file_hash
                _write_state(state)
            logger.info(
                "[Excel同步] 检测到已有同步文件，先将其作为基准；"
                "如需把文件内容导入数据库，请运行 sync_orders --import-now"
            )
        except Exception as exc:
            logger.warning(f"[Excel同步] 初始化读取文件失败：{exc}")
    else:
        export_sync_file()


def _watcher_loop() -> None:
    time.sleep(5)
    try:
        port = int(os.environ.get("ORDER_SYNC_LOCK_PORT", "8766"))
    except ValueError:
        port = 8766
    lock_sock = _try_bind_lock(port)
    if lock_sock is None:
        logger.info("[Excel同步] 其他进程已在运行同步 watcher，跳过")
        return
    try:
        interval = get_sync_interval()
        while True:
            try:
                if get_sync_enabled():
                    _ensure_initialized()
                    import_sync_file_if_changed()
            except Exception:
                logger.exception("[Excel同步] 同步循环发生异常")
            time.sleep(interval)
    finally:
        try:
            lock_sock.close()
        except Exception:
            pass


def start_sync_watcher():
    """应用启动时调用；设置 ORDER_SYNC_DISABLED=1 可关闭 watcher。"""
    if os.environ.get("ORDER_SYNC_DISABLED") == "1" or not get_sync_enabled():
        return None
    thread = threading.Thread(target=_watcher_loop, daemon=True, name="order-excel-sync")
    thread.start()
    return thread


def on_order_changed(sender, instance, **kwargs):
    """订单或出货记录变更时，安排一次导出。"""
    schedule_export()

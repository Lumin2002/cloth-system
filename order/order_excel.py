import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from io import BytesIO
from typing import Any, Callable, Optional

import pandas as pd
from django.db import transaction
import logging
logger = logging.getLogger(__name__)
from openpyxl import load_workbook

from .models import ClothOrder, Customer, Supplier
from .constants import DEFAULT_STAGES_MAP

SHEET_BULK = '大货'

# Excel 原始表头 -> 模型字段（与「大货」工作表一致）
EXCEL_COLUMN_MAP: dict[str, str] = {
    '序号': 'serial_number',
    '*下单日期': 'order_date',
    '*客户': 'customer',
    '跟单员': 'order_follower',
    '*订单类型': 'order_type',
    '*订单类似': 'order_type',
    '订单号': 'order_number',
    '款号': 'style_number',
    '规格': 'specification',
    '布种': 'cloth_type',
    '纺织类型': 'textile_type',
    '颜色': 'color',
    '色号': 'color_code',
    '成份': 'composition',
    '门幅': 'width',
    '克重': 'weight',
    '加工别': 'processing_type',
    '*订单数量': 'order_quantity',
    '单位': 'quantity_unit',
    '*价格': 'price',
    '单位.1': 'price_unit',
    '小缸费': 'small_vat_fee',
    '客人要求交期': 'customer_delivery_date',
    '大货进度跟踪': 'bulk_progress_tracking',
    '*成品出货对账总金额': 'finished_product_total_amount',
    '*付款方式\n（货前、货后）': 'payment_method',
    '付款方式\n（货前、货后）': 'supplier_payment_method',
    '*账期/天': 'payment_period',
    '*对账时间': 'reconciliation_date',
    '*货款支付时间': 'payment_date',
    '*是否付款': 'payment_status',
    '*是否逾期': 'overdue_status',
    '*是否开票': 'invoice_status',
    '*是否开证': 'certificate_status',
    '*证书类型': 'certificate_type',
    '最迟开证时间': 'latest_certificate_date',
    '实际操作时间': 'actual_operation_date',
    '实际开证时间': 'actual_certificate_date',
    '*成品供应商': 'finished_product_supplier',
    '档口': 'booth',
    '编号': 'supplier_code',
    '合同号': 'contract_number',
    '是否开票': 'supplier_invoice_status',
    '是否开票.1': 'supplier_invoice_status',
    '是否开证书': 'supplier_certificate_status',
    '对账时间': 'supplier_reconciliation_date',
    '货款支付时间': 'supplier_payment_date',
    '*已付款': 'supplier_paid',
    '*总金额': 'total_amount',
    '*成品成本价格': 'finished_product_cost_price',
    '单位.2': 'cost_price_unit',
    '*出货总数量': 'total_shipment_quantity',
    '单位.3': 'shipment_quantity_unit',
    '*成品出货时间': 'finished_product_shipment_date_1',
    '*成品出货数量': 'finished_product_shipment_quantity_1',
    '成品出货时间': 'finished_product_shipment_date_2',
    '成品出货数量': 'finished_product_shipment_quantity_2',
    '成品出货时间.1': 'finished_product_shipment_date_3',
    '成品出货数量.1': 'finished_product_shipment_quantity_3',
    '成品出货时间.2': 'finished_product_shipment_date_4',
    '成品出货数量.2': 'finished_product_shipment_quantity_4',
    '成品出货时间.3': 'finished_product_shipment_date_5',
    '成品出货数量.3': 'finished_product_shipment_quantity_5',
    '备注': 'remark',
    '地址': 'address',
}

DATE_FIELDS = {
    f for f in EXCEL_COLUMN_MAP.values()
    if f.endswith('_date') or 'date' in f
}
YES_NO_FIELDS = set()
BOOL_FIELDS = {'supplier_paid'}

DECIMAL_FIELDS = {
    'order_quantity', 'price', 'small_vat_fee', 'width', 'weight',
    'finished_product_total_amount', 'total_amount',
    'finished_product_cost_price', 'total_shipment_quantity',
    'finished_product_shipment_quantity_1', 'finished_product_shipment_quantity_2',
    'finished_product_shipment_quantity_3', 'finished_product_shipment_quantity_4',
    'finished_product_shipment_quantity_5',
}
INTEGER_FIELDS = {'serial_number', 'payment_period'}

MODEL_FIELD_NAMES = {
    f.name for f in ClothOrder._meta.get_fields()
    if hasattr(f, 'column') and not f.many_to_many
}

IMPORT_DEFAULTS = {
    'order_type': 'bulk',
    'payment_status': 'unpaid',
    'overdue_status': 'not_overdue',
    'invoice_status': 'not_invoiced',
    'certificate_status': 'not_invoiced',
    'payment_method': 'before_delivery',
    'certificate_type': 'none',
    'bulk_progress_tracking': '',
    'order_status': 'active',
}

EXPORT_COLUMNS: list[tuple[str, str]] = [
    ('序号', 'serial_number'),
    ('订单状态', 'order_status'),
    ('*下单日期', 'order_date'),
    ('*客户', 'customer'),
    ('跟单员', 'order_follower'),
    ('*订单类型', 'order_type'),
    ('订单号', 'order_number'),
    ('款号', 'style_number'),
    ('布种', 'cloth_type'),
    ('颜色', 'color'),
    ('*订单数量', 'order_quantity'),
    ('单位', 'quantity_unit'),
    ('*价格', 'price'),
    ('单位.1', 'price_unit'),
    ('*成品出货对账总金额', 'finished_product_total_amount'),
    ('已付款给供应商', 'supplier_paid'),
    ('*总金额', 'total_amount'),
    ('*是否付款', 'payment_status'),
    ('*是否逾期', 'overdue_status'),
    ('*成品供应商', 'finished_product_supplier'),
    ('*付款方式（货前、货后）', 'payment_method'),
    ('*对账时间', 'reconciliation_date'),
    ('*货款支付时间', 'payment_date'),
    ('备注', 'remark'),
    ('地址', 'address'),
    ('供应商已出货', 'supplier_shipped'),
]


def _normalize_header(name: Any) -> str:
    text = str(name).replace('*', '').replace('\n', ' ').strip()
    return re.sub(r'\s+', ' ', text)


def _is_serial_header(name: str) -> bool:
    return _normalize_header(name) == '序号'


def map_dataframe_columns(df: pd.DataFrame) -> pd.DataFrame:
    """将 Excel 原始列名映射为模型字段名（先精确匹配，再归一化匹配）。"""
    normalized_map: dict[str, str] = {}
    for excel_col, model_field in EXCEL_COLUMN_MAP.items():
        if model_field in MODEL_FIELD_NAMES:
            normalized_map[_normalize_header(excel_col)] = model_field

    rename_map: dict[str, str] = {}
    for col in df.columns:
        col_str = str(col).strip()
        if col_str in EXCEL_COLUMN_MAP and EXCEL_COLUMN_MAP[col_str] in MODEL_FIELD_NAMES:
            rename_map[col] = EXCEL_COLUMN_MAP[col_str]
            continue
        norm = _normalize_header(col_str)
        if norm in normalized_map:
            rename_map[col] = normalized_map[norm]
    return df.rename(columns=rename_map)


def _cell_has_strikethrough(cell) -> bool:
    """检测单元格字体是否带删除线。"""
    if cell is None or cell.value is None:
        return False
    font = cell.font
    if font is None:
        return False
    return bool(getattr(font, 'strike', None)) or bool(getattr(font, 'strikethrough', None))


def _row_has_data_from_cells(cells_by_col: dict[int, Any]) -> bool:
    for cell in cells_by_col.values():
        if cell.value is not None and str(cell.value).strip() != '':
            return True
    return False


def _row_is_cancelled_from_cells(
    cells_by_col: dict[int, Any],
    strike_check_cols: set[int],
    serial_col: Optional[int],
) -> bool:
    """
    汇总表约定：行内带删除线表示已取消订单。
    仅检查序号列及前几列有内容单元格，避免全表扫描过慢。
    """
    if serial_col and serial_col in cells_by_col:
        if _cell_has_strikethrough(cells_by_col[serial_col]):
            return True

    struck = 0
    filled = 0
    for col in strike_check_cols:
        cell = cells_by_col.get(col)
        if cell is None:
            continue
        if cell.value is None or str(cell.value).strip() == '':
            continue
        filled += 1
        if _cell_has_strikethrough(cell):
            struck += 1

    if filled == 0:
        return False
    return struck >= max(1, (filled + 1) // 2)


def read_bulk_sheet_with_status(
    file_obj,
    progress_callback: Optional[Callable[..., None]] = None,
) -> tuple[pd.DataFrame, list[bool]]:
    """
    用 openpyxl 读取「大货」表：返回 DataFrame 及每行是否已取消（删除线）。
    使用 iter_rows 单次遍历（read_only + cell() 随机访问会导致极慢甚至假死）。
    """
    raw = file_obj.read()
    file_obj.seek(0)

    _report_progress(progress_callback, 3, '正在打开 Excel 文件…', 'read')

    try:
        wb = load_workbook(filename=BytesIO(raw), data_only=True, read_only=False)
    except Exception as exc:
        raise ValueError('请上传 .xlsx 格式文件（需读取删除线格式，不支持旧版 .xls）') from exc

    sheet_name = SHEET_BULK if SHEET_BULK in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    _report_progress(progress_callback, 5, '正在读取表头…', 'read')

    header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=False))
    headers: dict[int, str] = {}
    seen_headers: dict[str, int] = {}
    for col_idx, cell in enumerate(header_cells, start=1):
        if cell.value is not None and str(cell.value).strip():
            raw = str(cell.value).strip()
            # 处理重复列名（如单位、单位.1）→与pandas读取时一致
            if raw in seen_headers:
                seen_headers[raw] += 1
                raw = f"{raw}.{seen_headers[raw]}"
            else:
                seen_headers[raw] = 0
            headers[col_idx] = raw

    if not headers:
        wb.close()
        return pd.DataFrame(), []

    min_col = min(headers.keys())
    max_col = max(headers.keys())
    serial_col = next((c for c, h in headers.items() if _is_serial_header(h)), None)
    strike_check_cols = set(list(headers.keys())[:15])
    if serial_col:
        strike_check_cols.add(serial_col)

    rows: list[dict[str, Any]] = []
    cancelled_flags: list[bool] = []
    empty_streak = 0
    max_empty_streak = 80
    scanned = 0
    sheet_max_row = ws.max_row or 1
    estimated = max(sheet_max_row - 1, 1)

    _report_progress(progress_callback, 6, '正在扫描数据行（检测删除线）…', 'read', 0, estimated)

    for row_cells in ws.iter_rows(
        min_row=2,
        min_col=min_col,
        max_col=max_col,
        values_only=False,
    ):
        scanned += 1
        if scanned % 100 == 0:
            pct = 6 + min(6, int(6 * scanned / estimated))
            _report_progress(
                progress_callback, pct,
                f'已扫描 {scanned} 行，有效 {len(rows)} 条…',
                'read', scanned, estimated,
            )

        cells_by_col: dict[int, Any] = {}
        for offset, cell in enumerate(row_cells):
            col_idx = min_col + offset
            if col_idx in headers:
                cells_by_col[col_idx] = cell

        if not _row_has_data_from_cells(cells_by_col):
            empty_streak += 1
            if empty_streak >= max_empty_streak:
                break
            continue
        empty_streak = 0

        cancelled_flags.append(
            _row_is_cancelled_from_cells(cells_by_col, strike_check_cols, serial_col)
        )
        rows.append({headers[col]: cells_by_col[col].value for col in headers if col in cells_by_col})

    wb.close()
    _report_progress(
        progress_callback, 12,
        f'解析完成，共 {len(rows)} 条有效数据',
        'read', len(rows), len(rows),
    )
    return pd.DataFrame(rows), cancelled_flags


def _scalar(value: Any) -> Any:
    if isinstance(value, pd.Series):
        return value.iloc[0] if len(value) else None
    return value


def _parse_date(value: Any):
    value = _scalar(value)
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d', '%Y-%m', '%Y/%m'):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        try:
            serial = float(text)
            return (datetime(1899, 12, 30) + timedelta(days=serial)).date()
        except ValueError:
            return None
    try:
        parsed = pd.to_datetime(value, errors='coerce')
        if pd.notna(parsed):
            return parsed.date()
    except (TypeError, ValueError):
        pass
    return None


def _parse_decimal(value: Any):
    value = _scalar(value)
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if pd.isna(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        text = str(value).strip().replace(',', '')
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None


def _parse_int(value: Any) -> Optional[int]:
    num = _parse_decimal(value)
    if num is None:
        return None
    return int(num)


def _parse_yes_no(value: Any) -> Optional[str]:
    text = str(_scalar(value) or '').strip().lower()
    if text in {'是', 'yes', 'true', 'y', '1', '已付款', '已付'}:
        return 'yes'
    if text in {'否', 'no', 'false', 'n', '0', '未付款', '未付'}:
        return 'no'
    return None


def _parse_bool(value: Any) -> Optional[bool]:
    """解析布尔值，支持Excel中的是/否、yes/no等"""
    text = str(_scalar(value) or '').strip().lower()
    if text in {'是', 'yes', 'true', 'y', '1', '已付款', '已付'}:
        return True
    if text in {'否', 'no', 'false', 'n', '0', '未付款', '未付'}:
        return False
    return None


def _parse_yes_no_status(value: Any, field_name: str) -> str:
    text = str(_scalar(value) or "").strip()

    _PAYMENT = {  # payment_status, supplier_payment_method
        "paid": {"paid", "已付款", "已付", "付清", "已付清", "yes", "true", "是", "y", "1"},
        "unpaid": {"unpaid", "未付款", "未付", "未付清", "no", "false", "否", "n", "0"},
        "partial": {"partial", "部分付款", "部分"},
    }
    _OVERDUE = {  # overdue_status
        "overdue": {"overdue", "逾期", "过期", "超期", "已逾期"},
        "not_overdue": {"not_overdue", "未逾期", "未过期", "正常"},
    }
    _INVOICE = {  # invoice_status, certificate_status, supplier_*
        "invoiced": {"invoiced", "已开票", "已开证", "已开证书", "已开", "是", "yes", "y", "1"},
        "not_invoiced": {"not_invoiced", "未开票", "未开证", "未开证书", "未开", "否", "no", "n", "0"},
    }

    mapping = _OVERDUE if field_name == "overdue_status" else _PAYMENT if "payment" in field_name else _INVOICE

    low = text.lower()
    for result_key, aliases in mapping.items():
        for alias in aliases:
            if text == alias or low == alias.lower():
                return result_key

    return "not_overdue" if field_name == "overdue_status" else "unpaid" if "payment" in field_name else "not_invoiced"


def _parse_choice_field(field_name: str, value: Any) -> Any:
    text = str(_scalar(value) or '').strip()
    if not text:
        return None
    if field_name == 'order_type':
        if '印花' in text:
            return 'bulk_print'
        if '大货订单' in text:
            return 'bulk'
        if '样板单' in text:
            return 'sample'
        return 'other'
    if field_name in ('payment_method', 'supplier_payment_method'):
        if '货后' in text:
            return 'after_delivery'
        if '货前' in text:
            return 'before_delivery'
        return 'before_delivery'
    if field_name == 'certificate_type':
        upper = text.upper()
        if 'GRS' in upper:
            return 'grs'
        if 'OEKO' in upper or 'TEX' in upper:
            return 'oeko'
        if '无' in text:
            return 'none'
        return 'other'
    if 'status' in field_name and field_name != 'order_status':
        return _parse_yes_no_status(value, field_name)
    return text


def parse_row(row: pd.Series) -> dict[str, Any]:
    """将一行 DataFrame 数据解析为模型字段字典。"""
    data: dict[str, Any] = {}
    for col in row.index:
        if col not in MODEL_FIELD_NAMES:
            continue
        raw = row[col]
        if pd.isna(_scalar(raw)):
            continue
        if col in DATE_FIELDS:
            parsed = _parse_date(raw)
            if parsed is not None:
                data[col] = parsed
        elif col in YES_NO_FIELDS:
            parsed = _parse_yes_no(raw)
            if parsed is not None:
                data[col] = parsed
        elif col in BOOL_FIELDS:
            parsed = _parse_bool(raw)
            if parsed is not None:
                data[col] = parsed
        elif col in DECIMAL_FIELDS:
            parsed = _parse_decimal(raw)
            if parsed is not None:
                data[col] = parsed
        elif col in INTEGER_FIELDS:
            parsed = _parse_int(raw)
            if parsed is not None:
                data[col] = parsed
        elif col in (
            'order_type', 'payment_method', 'supplier_payment_method',
            'certificate_type', 'payment_status', 'overdue_status',
            'invoice_status', 'certificate_status',
            'supplier_invoice_status', 'supplier_certificate_status',
        ):
            parsed = _parse_choice_field(col, raw)
            if parsed is not None:
                data[col] = parsed
        else:
            data[col] = str(_scalar(raw)).strip()

    # 单位与价格单位自动关联
    if "quantity_unit" in data and data["quantity_unit"] and "price_unit" not in data:
        u = data["quantity_unit"].strip()
        if u == "米":
            data["price_unit"] = "元/米"
        elif u == "码":
            data["price_unit"] = "元/码"

    # 出货总数量 = 5个分批次出货数量自动求和
    shipment_fields = [
        'finished_product_shipment_quantity_1',
        'finished_product_shipment_quantity_2',
        'finished_product_shipment_quantity_3',
        'finished_product_shipment_quantity_4',
        'finished_product_shipment_quantity_5',
    ]
    total = sum(
        float(data[f]) for f in shipment_fields
        if f in data and data[f] is not None
    )
    if total > 0:
        data['total_shipment_quantity'] = total

    return data


@dataclass
class ImportResult:
    imported: int = 0
    updated: int = 0
    skipped: int = 0
    cancelled: int = 0
    reactivated: int = 0
    errors: int = 0
    error_messages: list[str] = field(default_factory=list)

    @property
    def success(self) -> bool:
        return self.errors == 0

    def to_dict(self) -> dict[str, Any]:
        return {
            'imported': self.imported,
            'updated': self.updated,
            'skipped': self.skipped,
            'cancelled': self.cancelled,
            'reactivated': self.reactivated,
            'errors': self.errors,
            'error_messages': self.error_messages[:30],
            'success': self.success,
        }


def _report_progress(
    callback: Optional[Callable[..., None]],
    percent: int,
    message: str,
    phase: str = 'import',
    current: int = 0,
    total: int = 0,
) -> None:
    if callback:
        callback(
            percent=min(100, max(0, percent)),
            message=message,
            phase=phase,
            current=current,
            total=total,
        )


def import_orders_from_file(
    file_obj,
    progress_callback: Optional[Callable[..., None]] = None,
) -> ImportResult:
    """
    从 Excel 导入/更新订单（按序号 upsert）。
    带删除线的行标记为 order_status=cancelled，否则为 active。
    """
    df, cancelled_flags = read_bulk_sheet_with_status(file_obj, progress_callback=progress_callback)
    df = map_dataframe_columns(df)
    result = ImportResult()

    total_rows = len(df)
    if total_rows == 0:
        _report_progress(progress_callback, 100, '工作表中没有可导入的数据行', 'done', 0, 0)
        return result

    _report_progress(
        progress_callback, 12,
        f'共 {total_rows} 条，正在写入数据库…',
        'import', 0, total_rows,
    )

    with transaction.atomic():
        for i, (index, row) in enumerate(df.iterrows()):
            if i % 5 == 0 or i == total_rows - 1:
                pct = 12 + int(83 * (i + 1) / total_rows)
                _report_progress(
                    progress_callback, pct,
                    f'正在处理 {i + 1} / {total_rows} 条…',
                    'import', i + 1, total_rows,
                )
            try:
                row_data = parse_row(row)
                serial_number = row_data.get('serial_number')
                if serial_number is None:
                    result.skipped += 1
                    continue
                # 缺少下单日期的行不导入
                if not row_data.get('order_date'):
                    result.skipped += 1
                    continue

                # 检测是否已有出货数据，有则标记供应商已出货
                if row_data.get('total_amount') and row_data.get('total_shipment_quantity') and row_data.get('finished_product_cost_price'):
                    row_data['supplier_shipped'] = True

                is_cancelled = cancelled_flags[i] if i < len(cancelled_flags) else False
                row_data['order_status'] = 'cancelled' if is_cancelled else 'active'

                for key, default in IMPORT_DEFAULTS.items():
                    if key != 'order_status':
                        row_data.setdefault(key, default)

                _ensure_customer(row_data.get("customer"))
                existing = ClothOrder.objects.filter(serial_number=serial_number).first()
                if existing:
                    was_cancelled = existing.order_status == 'cancelled'
                    for key, value in row_data.items():
                        if value is not None and key != 'serial_number':
                            setattr(existing, key, value)
                    existing.save()
                    _sync_shipments_from_order(existing)
                    _match_supplier(existing)
                    result.updated += 1
                    if is_cancelled and not was_cancelled:
                        result.cancelled += 1
                    elif not is_cancelled and was_cancelled:
                        result.reactivated += 1
                else:
                    new_order = ClothOrder.objects.create(**row_data)
                    _sync_shipments_from_order(new_order)
                    _match_supplier(new_order)
                    result.imported += 1
                    if is_cancelled:
                        result.cancelled += 1
            except Exception as exc:
                result.errors += 1
                result.error_messages.append(f'第 {index + 2} 行: {exc}')

    _report_progress(progress_callback, 100, '导入完成', 'done', total_rows, total_rows)
    fix_imported_progress_stages()
    return result


def _match_supplier(order):
    if order.supplier_id or not order.finished_product_supplier:
        return
    try:
        supplier = Supplier.objects.filter(
            company_name__iexact=order.finished_product_supplier.strip()
        ).first()
        if supplier:
            order.supplier = supplier
            order.save(update_fields=["supplier"])
    except Exception:
        pass


def _ensure_customer(name):
    """订单导入时自动补充不存在的客户。"""
    name = str(name or "").strip()
    if name:
        Customer.objects.get_or_create(name=name)


def _sync_shipments_from_order(order):
    """从订单的旧批次字段同步出货记录到 Shipment 表"""
    from .models import Shipment
    for i in range(1, 6):
        date = getattr(order, f'finished_product_shipment_date_{i}', None)
        qty = getattr(order, f'finished_product_shipment_quantity_{i}', None)
        if date or qty:
            Shipment.objects.get_or_create(
                order=order,
                batch_number=i,
                defaults={"date": date, "quantity": qty},
            )


def fix_imported_progress_stages():
    """导入后将订单进度阶段与订单类型匹配"""
    for order_type, stages in DEFAULT_STAGES_MAP.items():
        count = ClothOrder.objects.filter(order_type=order_type).exclude(progress_stages=stages).update(progress_stages=stages)
        if count:
            logger.info(f"  {order_type}: {count} 条更新")


def _display_value(order: ClothOrder, field_name: str):
    value = getattr(order, field_name, None)
    if value is None:
        return None
    if field_name == 'order_status':
        return order.get_order_status_display()
    if field_name == 'order_type':
        return order.get_order_type_display()
    if field_name == 'payment_status':
        return order.get_payment_status_display()
    if field_name == 'overdue_status':
        return order.get_overdue_status_display()
    if field_name == 'payment_method':
        return order.get_payment_method_display()
    if field_name == 'paid_amount':
        if value == 'yes':
            return '是'
        if value == 'no':
            return '否'
        return None
    if hasattr(value, 'isoformat'):
        return value
    if isinstance(value, (int, float)):
        return value
    return str(value)


def export_orders_dataframe(queryset=None) -> pd.DataFrame:
    """将订单 queryset 转为与汇总表一致的 DataFrame。"""
    if queryset is None:
        queryset = ClothOrder.objects.all().order_by('-order_date', '-serial_number')
    rows = []
    for order in queryset:
        rows.append({
            excel_col: _display_value(order, model_field)
            for excel_col, model_field in EXPORT_COLUMNS
        })
    return pd.DataFrame(rows)

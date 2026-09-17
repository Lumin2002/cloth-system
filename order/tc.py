import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .models import ClothOrder


TC_HEADERS = [
    "序号",
    "客户",
    "跟单员",
    "布种",
    "成分",
    "门幅",
    "克重",
    "颜色",
    "规格",
    "订单数量",
    "单位",
    "价格",
    "单位",
    "成品出货时间",
    "成品出货数量",
    "金额",
    "订单号",
    "款号",
    "件数",
    "",
    "开票账户",
    "对应重量（kg）",
    "再生纤维重量（kg）",
]

TC_COLUMN_WIDTHS = [
    5.43,
    7.57,
    7.86,
    12.0,
    28.29,
    6.43,
    8.43,
    12.29,
    6.43,
    6.29,
    4.43,
    8.0,
    6.43,
    13.0,
    9.14,
    10.86,
    25.14,
    39.29,
    44.86,
    8.43,
    9.14,
    8.29,
    8.29,
]

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

FONT_HEADER = Font(name="Arial", size=9, bold=True)
FONT_HEADER_LARGE = Font(name="Arial", size=11, bold=True)
FONT_HEADER_SONG = Font(name="宋体", size=10, bold=True)
FONT_HEADER_SONG_SMALL = Font(name="宋体", size=9, bold=True)
FONT_HEADER_ACCOUNT = Font(name="Arial", size=10)
FONT_BODY = Font(name="Arial", size=10)
FONT_BODY_SMALL = Font(name="Arial", size=8)
FONT_BODY_BOLD = Font(name="Arial", size=11, bold=True)
FONT_BODY_LARGE = Font(name="Arial", size=11)
FONT_BODY_SONG = Font(name="宋体", size=10)


def _font_for_header(col: int) -> Font:
    if col in (4, 14):
        return FONT_HEADER_LARGE
    if col in (17, 18):
        return FONT_HEADER_SONG
    if col == 19:
        return FONT_HEADER_SONG_SMALL
    if col in (21, 22, 23):
        return FONT_HEADER_ACCOUNT
    return FONT_HEADER


def _setup_sheet(ws):
    for index, width in enumerate(TC_COLUMN_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_setup.scale = 75

    for col, header in enumerate(TC_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = _font_for_header(col)
        cell.alignment = ALIGN_CENTER_WRAP if col in (17, 18) else ALIGN_LEFT_WRAP
        cell.border = THIN_BORDER


def _set_cell(
    ws,
    row,
    col,
    value=None,
    font=FONT_BODY,
    alignment=ALIGN_LEFT,
    number_format=None,
):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    cell.alignment = alignment
    cell.border = THIN_BORDER
    if number_format:
        cell.number_format = number_format
    return cell


def _textile_type_for_tc(order: ClothOrder) -> str:
    display = order.get_textile_type_display() or order.textile_type or ""
    if display.endswith("布"):
        return display[:-1]
    return display or (order.specification or "")


def _add_order_rows(ws, order: ClothOrder, start_row: int) -> int:
    row = start_row
    for shipment in order.shipments.filter(is_deleted=False).order_by("batch_number"):
        shipment_date = shipment.date
        shipment_qty = shipment.quantity
        if shipment_date is None and shipment_qty is None:
            continue

        qty = float(shipment_qty) if shipment_qty is not None else 0.0
        price = float(order.price) if order.price is not None else 0.0
        width = float(order.width) if order.width is not None else 0.0
        weight = float(order.weight) if order.weight is not None else 0.0

        _set_cell(ws, row, 1, order.serial_number or "")
        _set_cell(ws, row, 2, order.customer or "", font=FONT_BODY_SMALL)
        _set_cell(
            ws,
            row,
            3,
            order.order_follower or "",
            font=FONT_BODY_SMALL,
            alignment=ALIGN_LEFT_WRAP,
        )
        _set_cell(ws, row, 4, order.cloth_type or "", font=FONT_BODY_BOLD)
        _set_cell(ws, row, 5, order.composition or "")
        _set_cell(ws, row, 6, width if order.width is not None else "")
        _set_cell(ws, row, 7, weight if order.weight is not None else "")
        _set_cell(ws, row, 8, order.color or "")
        _set_cell(ws, row, 9, _textile_type_for_tc(order), font=FONT_BODY_SONG)
        _set_cell(ws, row, 10, float(order.order_quantity) if order.order_quantity is not None else "")
        _set_cell(
            ws,
            row,
            11,
            order.quantity_unit or "米",
            font=FONT_BODY_SMALL,
            alignment=ALIGN_LEFT_WRAP,
        )
        _set_cell(
            ws,
            row,
            12,
            price if order.price is not None else "",
            number_format='"¥"#,##0.00;[Red]"¥"\\-#,##0.00',
        )
        _set_cell(
            ws,
            row,
            13,
            order.price_unit or "元/米",
            font=FONT_BODY_SMALL,
            alignment=ALIGN_LEFT_WRAP,
        )
        _set_cell(
            ws,
            row,
            14,
            shipment_date,
            font=FONT_BODY_LARGE,
            alignment=ALIGN_LEFT_WRAP,
            number_format="yyyy/m/d;@",
        )
        _set_cell(ws, row, 15, qty if shipment_qty is not None else "")
        _set_cell(ws, row, 16, f"=O{row}*L{row}", number_format="#,##0.00")
        _set_cell(ws, row, 17, order.order_number or "", alignment=ALIGN_CENTER)
        _set_cell(ws, row, 18, order.style_number or "", alignment=ALIGN_CENTER_WRAP)
        _set_cell(ws, row, 19, "")
        _set_cell(ws, row, 20, "")
        _set_cell(
            ws,
            row,
            21,
            order.customer or "",
            font=FONT_BODY_SONG,
            alignment=ALIGN_LEFT_WRAP,
        )
        _set_cell(
            ws,
            row,
            22,
            f"=G{row}/1000*F{row}/100*O{row}",
            font=FONT_BODY_SMALL,
            alignment=ALIGN_LEFT_WRAP,
            number_format="0.0",
        )
        _set_cell(
            ws,
            row,
            23,
            f"=V{row}*1",
            font=FONT_BODY_SMALL,
            alignment=ALIGN_LEFT_WRAP,
            number_format="0.0",
        )
        ws.row_dimensions[row].height = 20
        row += 1
    return row


def generate_bulk_tc(orders):
    wb = Workbook()
    ws = wb.active
    ws.title = "TC表"
    _setup_sheet(ws)

    ordered_orders = sorted(
        orders,
        key=lambda order: (
            order.serial_number is None,
            order.serial_number or 0,
        ),
    )
    row = 2
    for order in ordered_orders:
        row = _add_order_rows(ws, order, row)

    if row == 2:
        _set_cell(
            ws,
            row,
            1,
            "暂无出货记录",
            alignment=ALIGN_LEFT_WRAP,
        )
        row += 1

    ws.auto_filter.ref = f"A1:W{row - 1}"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

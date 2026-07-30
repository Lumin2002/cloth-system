import openpyxl
import re
from datetime import datetime
from django.core.management.base import BaseCommand
from order.models import FabricQuotation


def parse_price(text):
    """解析价格文本，返回 (solid_price, print_price, unit)"""
    if not text or not isinstance(text, str):
        return None, None, "M"
    text = text.strip()
    if text == "" or text == "N/A":
        return None, None, "M"

    # 去掉开头的 RMB 或 RMB
    t = re.sub(r'^(RMB|RMB)\s*', '', text, flags=re.IGNORECASE).strip()

    # 判断单位
    unit = "M"
    if "/kg" in t.lower() or "/KG" in t:
        unit = "KG"

    # 处理逗号分隔的： "13/M for solid, RMB 15/m for print"
    solid_val = None
    print_val = None

    parts = re.split(r'[,，]', t)
    for part in parts:
        part = part.strip()
        part = re.sub(r'^(RMB|RMB)\s*', '', part, flags=re.IGNORECASE).strip()
        is_print = bool(re.search(r'for\s+print', part, re.IGNORECASE))
        is_solid = bool(re.search(r'for\s+solid|for\s+soild', part, re.IGNORECASE)) or not is_print

        # 提取数字
        m = re.search(r'(\d+(?:\.\d+)?)', part)
        val = float(m.group(1)) if m else None

        if is_print and val is not None:
            print_val = val
        if is_solid and val is not None:
            solid_val = val

    # 如果没分开（单个数字），尝试提取
    if solid_val is None and print_val is None:
        m = re.search(r'(\d+(?:\.\d+)?)', t)
        if m:
            solid_val = float(m.group(1))

    # 检查是否同时有 solid 和 print
    if solid_val is not None and print_val is not None:
        return solid_val, print_val, unit
    return solid_val, print_val, unit


class Command(BaseCommand):
    help = "从 GFG Fabric Quotation Chart.xlsx 导入报价数据"

    def add_arguments(self, parser):
        parser.add_argument("file", nargs="?", default="GFG Fabric Quotation Chart from20250416.xlsx")

    def handle(self, *args, **options):
        path = options["file"]
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Summary"]

        FabricQuotation.objects.all().delete()
        created = 0
        skipped = 0

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True), 2):
            # Excel 第3行是填写说明，跳过
            if row_idx == 3:
                continue
            # A列数字
            a_val = row[0]
            # H列（第8列）- 面料编号，没有就跳过
            article_no = str(row[7]).strip() if row[7] else ""
            if not article_no or article_no in ("填面料编号", "N/A"):
                skipped += 1
                continue

            # 日期
            date_sent = None
            if isinstance(row[1], datetime):
                date_sent = row[1].date()
            elif hasattr(row[1], "date"):
                date_sent = row[1].date()

            # 价格解析
            price_200m_text = str(row[11]).strip() if row[11] else ""
            p200m, p200m_print, p200m_unit = parse_price(price_200m_text)

            price_regular_text = str(row[13]).strip() if row[13] else ""
            preg, preg_print, preg_unit = parse_price(price_regular_text)

            # 克重和门幅
            weight = str(row[9]).strip() if row[9] else ""
            cuttable_width = str(row[10]).strip() if row[10] else ""

            # 备注
            remark = str(row[18]).strip() if row[18] else ""

            # 后面几列（U列之后）
            extra_remark = ""
            for col_idx in [20, 21, 22, 23, 24, 25, 26, 27, 28, 29]:
                if col_idx < len(row) and row[col_idx] is not None:
                    v = str(row[col_idx]).strip()
                    if v:
                        extra_remark += v + " "
            extra_remark = extra_remark.strip()

            FabricQuotation.objects.create(
                row_index=int(a_val) if isinstance(a_val, (int, float)) else None,
                date_sent=date_sent,
                supplier_name=str(row[2]).strip() if row[2] else "",
                supplier_contact=str(row[3]).strip() if row[3] else "",
                gfg_dev_no=str(row[4]).strip() if row[4] else "",
                preferred_material=str(row[5]).strip() if row[5] else "",
                color_card=str(row[6]).strip() if row[6] else "",
                article_no=article_no.replace("（换成", "（换成"),
                composition=str(row[8]).strip() if row[8] else "",
                weight=weight,
                cuttable_width=cuttable_width,
                price_200m_text=price_200m_text,
                price_200m=p200m,
                price_200m_print=p200m_print,
                price_200m_unit=p200m_unit,
                regular_mcq=str(row[12]).strip() if row[12] else "",
                price_regular_text=price_regular_text,
                price_regular=preg,
                price_regular_print=preg_print,
                price_regular_unit=preg_unit,
                lead_time=str(row[14]).strip() if row[14] else "",
                price_validity=str(row[15]).strip() if row[15] else "",
                yarn_count=str(row[16]).strip() if row[16] else "",
                density_gauge=str(row[17]).strip() if row[17] else "",
                remark=remark,
                quoted_to=str(row[19]).strip() if len(row) > 19 and row[19] else "",
                extra_remark=extra_remark,
            )
            created += 1

        self.stdout.write(self.style.SUCCESS(f"导入完成: 创建 {created} 条, 跳过 {skipped} 条"))

"""视图模块：面料报价（由 views.py 拆分而来）。"""

import re
from datetime import date, datetime
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import transaction
from django.db.models import (
    Count,
    DateField,
    DecimalField,
    Exists,
    ExpressionWrapper,
    F,
    FloatField,
    IntegerField,
    OuterRef,
    Prefetch,
    Q,
    Subquery,
    Sum,
    Value,
)
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_GET, require_POST
from django.views.generic import CreateView, DetailView, ListView, UpdateView
from .decorators import admin_required, validate_file_upload
from .i18n import t
from .models import (
    ClothCatalog,
    ClothOrder,
    FabricQuotation,
    InventoryItem,
    InventoryLog,
    Notification,
    Shipment,
    Supplier,
    notify_all_staff,
    notify_user,
)
from .views_common import AdminRequiredMixin, logger


@admin_required
def quotation_lookup(request):
    article_no = request.GET.get("q", "").strip()
    if not article_no:
        return JsonResponse({"error": "Need q parameter"}, status=400)

    qs = FabricQuotation.objects.filter(article_no__icontains=article_no).order_by("-date_sent")
    if not qs.exists():
        return JsonResponse({"found": False})

    q = qs.first()
    data = {
        "found": True,
        "id": q.id,
        "supplier_name": q.supplier_name,
        "article_no": q.article_no,
        "composition": q.composition,
        "weight": q.weight,
        "cuttable_width": q.cuttable_width,
        "price_200m_text": q.price_200m_text,
        "price_200m": str(q.price_200m) if q.price_200m else None,
        "price_200m_print": str(q.price_200m_print) if q.price_200m_print else None,
        "price_200m_unit": q.price_200m_unit,
        "regular_mcq": q.regular_mcq,
        "price_regular_text": q.price_regular_text,
        "price_regular": str(q.price_regular) if q.price_regular else None,
        "price_regular_print": str(q.price_regular_print) if q.price_regular_print else None,
        "price_regular_unit": q.price_regular_unit,
        "lead_time": q.lead_time,
        "remark": q.remark,
        "preferred_material": q.preferred_material,
        "quoted_to": q.quoted_to,
        "date_sent": str(q.date_sent) if q.date_sent else None,
    }
    return JsonResponse(data)


class QuotationListView(AdminRequiredMixin, LoginRequiredMixin, ListView):
    model = FabricQuotation
    template_name = "order/quotation_list.html"
    context_object_name = "items"
    paginate_by = 50

    def get_queryset(self):
        qs = FabricQuotation.objects.all()
        q = self.request.GET.get("q", "").strip()
        if q:
            qs = qs.filter(
                Q(article_no__icontains=q)
                | Q(composition__icontains=q)
                | Q(supplier_name__icontains=q)
            )
        return qs

    def get_paginate_by(self, queryset):
        try:
            return int(self.request.GET.get("per_page", 50))
        except (ValueError, TypeError):
            return 50

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["q"] = self.request.GET.get("q", "")
        ctx["total"] = FabricQuotation.objects.count()
        ctx["per_page_options"] = [20, 50, 100]
        ctx["per_page"] = self.request.GET.get("per_page", 50)
        try:
            ctx["per_page"] = int(ctx["per_page"])
        except ValueError:
            ctx["per_page"] = 50
        return ctx


class QuotationDetailView(AdminRequiredMixin, LoginRequiredMixin, DetailView):
    model = FabricQuotation
    template_name = "order/quotation_detail.html"
    context_object_name = "item"


class QuotationCreateView(AdminRequiredMixin, LoginRequiredMixin, CreateView):
    model = FabricQuotation
    fields = [
        "article_no",
        "supplier_name",
        "supplier_contact",
        "date_sent",
        "gfg_dev_no",
        "preferred_material",
        "color_card",
        "composition",
        "weight",
        "cuttable_width",
        "yarn_count",
        "density_gauge",
        "lead_time",
        "price_validity",
        "price_200m_text",
        "price_200m",
        "price_200m_print",
        "price_200m_unit",
        "regular_mcq",
        "price_regular_text",
        "price_regular",
        "price_regular_print",
        "price_regular_unit",
        "remark",
        "quoted_to",
        "extra_remark",
    ]
    template_name = "order/quotation_form.html"
    success_url = "/quotation/"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["form_title"] = t("label.quotation_add")
        return ctx

    def form_valid(self, form):
        response = super().form_valid(form)
        logger.info(
            t(
                "log.quotation_create",
                username=self.request.user.username,
                article=self.object.article_no,
            )
        )
        return response


@admin_required
def quotation_delete(request, pk):
    obj = get_object_or_404(FabricQuotation, pk=pk)
    article = obj.article_no
    obj.delete()
    logger.info(t("log.quotation_delete", username=request.user.username, article=article))
    messages.success(request, t("msg.quotation_deleted", article=article))
    return redirect("quotation_list")


def _parse_price(text):
    """解析价格文本，返回 (素色价格, 印花价格, 单位)"""
    if not text or not isinstance(text, str):
        return None, None, "M"
    text = text.strip()
    if not text or text == "N/A":
        return None, None, "M"

    cleaned = re.sub(r"^(RMB|RMB)\s*", "", text, flags=re.IGNORECASE).strip()
    unit = "KG" if "/kg" in cleaned.lower() else "M"

    solid_val = None
    print_val = None

    parts = re.split(r"[,，]", cleaned)
    for part in parts:
        part = part.strip()
        part = re.sub(r"^(RMB|RMB)\s*", "", part, flags=re.IGNORECASE).strip()
        is_print = bool(re.search(r"for\s+print", part, re.IGNORECASE))
        is_solid = bool(re.search(r"for\s+solid|for\s+soild", part, re.IGNORECASE)) or not is_print

        m = re.search(r"(\d+(?:\.\d+)?)", part)
        val = float(m.group(1)) if m else None

        if is_print and val is not None:
            print_val = val
        if is_solid and val is not None:
            solid_val = val

    if solid_val is None and print_val is None:
        m = re.search(r"(\d+(?:\.\d+)?)", cleaned)
        if m:
            solid_val = float(m.group(1))

    return solid_val, print_val, unit


@admin_required
def quotation_import(request):
    if request.method == "POST":
        excel_file = request.FILES.get("excel_file")
        if not excel_file:
            messages.error(request, t("msg.select_file"))
            return redirect("quotation_import")

        try:
            import openpyxl

            wb = openpyxl.load_workbook(excel_file, data_only=True)
            if "Summary" not in wb.sheetnames:
                messages.error(request, t("msg.excel_summary_missing"))
                return redirect("quotation_import")
            ws = wb["Summary"]
        except Exception as e:
            messages.error(request, t("msg.read_excel_error", e=e))
            return redirect("quotation_import")

        batch = []
        imported = 0
        skipped = 0

        try:
            with transaction.atomic():
                for row_idx, row in enumerate(
                    ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True), 2
                ):
                    if row_idx == 3:
                        continue

                    article_no = str(row[7]).strip() if row[7] else ""
                    if not article_no:
                        skipped += 1
                        continue

                    date_sent = None
                    if isinstance(row[1], datetime):
                        date_sent = row[1].date()

                    price_200m_text = str(row[11]).strip() if row[11] else ""
                    p200m, p200m_print, p200m_unit = _parse_price(price_200m_text)

                    price_regular_text = str(row[13]).strip() if row[13] else ""
                    preg, preg_print, preg_unit = _parse_price(price_regular_text)

                    extra = ""
                    for ci in [20, 21, 22, 23, 24, 25, 26, 27, 28, 29]:
                        if ci < len(row) and row[ci] is not None:
                            v = str(row[ci]).strip()
                            if v:
                                extra += v + " "
                    extra = extra.strip()

                    batch.append(
                        FabricQuotation(
                            row_index=int(row[0]) if isinstance(row[0], (int, float)) else None,
                            date_sent=date_sent,
                            supplier_name=str(row[2]).strip() if row[2] else "",
                            supplier_contact=str(row[3]).strip() if row[3] else "",
                            gfg_dev_no=str(row[4]).strip() if row[4] else "",
                            preferred_material=str(row[5]).strip() if row[5] else "",
                            color_card=str(row[6]).strip() if row[6] else "",
                            article_no=article_no,
                            composition=str(row[8]).strip() if row[8] else "",
                            weight=str(row[9]).strip() if row[9] else "",
                            cuttable_width=str(row[10]).strip() if row[10] else "",
                            price_200m_text=price_200m_text,
                            price_200m=p200m,
                            price_200m_print=p200m_print,
                            price_200m_unit=p200m_unit,
                            regular_mcq=str(row[12]).strip() if row[12] else "",
                            price_regular_text=price_regular_text,
                            price_regular=preg,
                            price_regular_print=preg_print,
                            price_regular_unit=preg_unit,
                            lead_time=str(row[14]).strip() if row[14] else "",
                            price_validity=str(row[15]).strip() if row[15] else "",
                            yarn_count=str(row[16]).strip() if row[16] else "",
                            density_gauge=str(row[17]).strip() if row[17] else "",
                            remark=str(row[18]).strip() if row[18] else "",
                            quoted_to=str(row[19]).strip() if len(row) > 19 and row[19] else "",
                            extra_remark=extra,
                        )
                    )
                    imported += 1

                # 事务内先删后插，失败自动回滚
                FabricQuotation.objects.all().delete()
                FabricQuotation.objects.bulk_create(batch)

            messages.success(request, t("msg.import_complete", imported=imported, skipped=skipped))
            logger.info(
                t(
                    "log.quotation_import_done",
                    username=request.user.username,
                    imported=imported,
                    skipped=skipped,
                )
            )
            return redirect("quotation_list")

        except Exception as e:
            messages.error(request, t("msg.import_failed", error=str(e)))
            logger.error(f"Quotation import failed: {e}", exc_info=True)
            return redirect("quotation_import")

    return render(request, "order/quotation_import.html")


@admin_required
def quotation_delete_all(request):
    count = FabricQuotation.objects.all().delete()[0]
    logger.info(t("log.quotation_delete_all", username=request.user.username, count=count))
    messages.success(request, t("msg.all_quotation_cleared", count=count))
    return redirect("quotation_list")

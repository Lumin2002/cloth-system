import json
import os
import subprocess
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path
from time import time
from unittest.mock import MagicMock, patch
from urllib.error import URLError

from openpyxl import load_workbook
from django.conf import settings
from django.contrib.auth.models import AnonymousUser, User
from django.core.cache import cache
from django.test import RequestFactory, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from order.captcha import CAPTCHA_SESSION_KEY
from order import service_monitor
from order import terminal
from order.forms import ShipmentForm
from order.middleware import SESSION_LAST_ACTIVITY_KEY
from order.models import ClothOrder, Notification, Shipment, Supplier
from order.views_common import handler400, handler403, handler404, handler500


@override_settings(
    MIDDLEWARE=[
        "django.middleware.security.SecurityMiddleware",
        "django.contrib.sessions.middleware.SessionMiddleware",
        "django.middleware.common.CommonMiddleware",
        "django.middleware.csrf.CsrfViewMiddleware",
        "django.contrib.auth.middleware.AuthenticationMiddleware",
        "django.contrib.messages.middleware.MessageMiddleware",
        "order.middleware.SessionTimeoutMiddleware",
        "django.middleware.clickjacking.XFrameOptionsMiddleware",
    ],
    SUPPLIER_SESSION_IDLE_TIMEOUT_MINUTES=30,
    SUPERUSER_SESSION_IDLE_TIMEOUT_MINUTES=15,
    SESSION_TIMEOUT_WARNING_SECONDS=60,
)
class SessionTimeoutMiddlewareTests(TestCase):
    def setUp(self):
        self.supplier_user = User.objects.create_user(
            username="supplier_timeout",
            email="supplier@example.com",
            password="test-pass-123",
        )
        self.supplier = Supplier.objects.create(
            user=self.supplier_user,
            company_name="超时测试供应商",
        )
        self.admin_user = User.objects.create_superuser(
            username="admin_timeout",
            email="admin@example.com",
            password="test-pass-123",
        )

    def _expire_session(self, minutes_ago=31):
        session = self.client.session
        session[SESSION_LAST_ACTIVITY_KEY] = time() - minutes_ago * 60
        session.save()

    def test_supplier_session_expires_after_idle_timeout(self):
        self.client.force_login(self.supplier_user)
        self._expire_session()

        response = self.client.get(reverse("supplier_dashboard"))

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("home"))
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_supplier_active_session_stays_logged_in(self):
        self.client.force_login(self.supplier_user)
        session = self.client.session
        session[SESSION_LAST_ACTIVITY_KEY] = time()
        session.save()

        response = self.client.get(reverse("supplier_dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertIn("_auth_user_id", self.client.session)
        self.assertLessEqual(self.client.session.get_expiry_age(), 30 * 60 + 1)

    def test_superuser_session_expires_after_idle_timeout(self):
        self.client.force_login(self.admin_user)
        self._expire_session()

        response = self.client.get(reverse("dashboard"))

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("home"))
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_superuser_active_session_stays_logged_in(self):
        self.client.force_login(self.admin_user)
        session = self.client.session
        session[SESSION_LAST_ACTIVITY_KEY] = time()
        session.save()

        response = self.client.get(reverse("dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertIn("_auth_user_id", self.client.session)

    def test_superuser_admin_request_redirects_to_admin_login(self):
        self.client.force_login(self.admin_user)
        self._expire_session()

        response = self.client.get("/admin/")

        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith("/admin/login/"))
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_background_poll_does_not_extend_session(self):
        self.client.force_login(self.supplier_user)
        old_ts = time() - 20 * 60
        session = self.client.session
        session[SESSION_LAST_ACTIVITY_KEY] = old_ts
        session.save()

        response = self.client.get("/notifications/unread-count/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.session.get(SESSION_LAST_ACTIVITY_KEY), old_ts)

    def test_expired_background_poll_returns_401(self):
        self.client.force_login(self.supplier_user)
        self._expire_session()

        response = self.client.get("/notifications/unread-count/")

        self.assertEqual(response.status_code, 401)
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_login_then_page_request_sets_session_activity_timestamp(self):
        session = self.client.session
        session[CAPTCHA_SESSION_KEY] = {"code": "TEST", "expires": time() + 300}
        session.save()
        self.client.post(
            reverse("home"),
            {
                "username": self.supplier_user.username,
                "password": "test-pass-123",
                "captcha": "test",
            },
        )
        self.client.get(reverse("supplier_dashboard"))
        self.assertIn(SESSION_LAST_ACTIVITY_KEY, self.client.session)


class ServiceMonitorTests(TestCase):
    def test_database_status_ok(self):
        cursor = MagicMock()
        cursor_cm = MagicMock()
        cursor_cm.__enter__.return_value = cursor
        conn = MagicMock()
        conn.vendor = "mysql"
        conn.cursor.return_value = cursor_cm

        with patch("order.service_monitor.connection", conn):
            result = service_monitor._database_status()

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["engine"], "mysql")

    def test_database_status_error(self):
        conn = MagicMock()
        conn.vendor = "mysql"
        conn.cursor.side_effect = Exception("db down")

        with patch("order.service_monitor.connection", conn):
            result = service_monitor._database_status()

        self.assertEqual(result["status"], "error")

    @override_settings(REDIS_URL="")
    def test_redis_not_configured(self):
        result = service_monitor._redis_status()
        self.assertEqual(result["status"], "not_configured")

    @override_settings(REDIS_URL="redis://127.0.0.1:6379/0")
    def test_redis_status_ok(self):
        client = MagicMock()
        client.connection_pool.connection_kwargs = {}
        client.ping.return_value = True

        with patch("order.service_monitor.get_redis_connection", return_value=client), patch.object(
            service_monitor, "DJANGO_REDIS_AVAILABLE", True
        ):
            result = service_monitor._redis_status()

        self.assertEqual(result["status"], "ok")

    @override_settings(REDIS_URL="redis://127.0.0.1:6379/0")
    def test_redis_status_error(self):
        client = MagicMock()
        client.connection_pool.connection_kwargs = {}
        client.ping.side_effect = Exception("redis down")

        with patch("order.service_monitor.get_redis_connection", return_value=client), patch.object(
            service_monitor, "DJANGO_REDIS_AVAILABLE", True
        ):
            result = service_monitor._redis_status()

        self.assertEqual(result["status"], "error")

    @override_settings(NGINX_CHECK_URL="")
    def test_nginx_not_configured(self):
        result = service_monitor._nginx_status()
        self.assertEqual(result["status"], "not_configured")

    @override_settings(NGINX_CHECK_URL="http://127.0.0.1:8233/")
    def test_nginx_status_ok(self):
        resp = MagicMock()
        resp.status = 200
        resp_cm = MagicMock()
        resp_cm.__enter__.return_value = resp

        with patch("order.service_monitor.urllib.request.urlopen", return_value=resp_cm) as urlopen:
            result = service_monitor._nginx_status()

        self.assertEqual(result["status"], "ok")
        urlopen.assert_called_once()
        request = urlopen.call_args.args[0]
        self.assertEqual(request.get_method(), "GET")

    @override_settings(NGINX_CHECK_URL="https://your-nginx.example.com:40614")
    def test_nginx_https_uses_unverified_context(self):
        resp = MagicMock()
        resp.status = 200
        resp_cm = MagicMock()
        resp_cm.__enter__.return_value = resp

        with patch("order.service_monitor.urllib.request.urlopen", return_value=resp_cm) as urlopen:
            result = service_monitor._nginx_status()

        self.assertEqual(result["status"], "ok")
        self.assertIsNotNone(urlopen.call_args.kwargs.get("context"))
        request = urlopen.call_args.args[0]
        self.assertEqual(request.get_method(), "GET")

    @override_settings(NGINX_CHECK_URL="http://127.0.0.1:8233/")
    def test_nginx_status_error(self):
        with patch("order.service_monitor.urllib.request.urlopen", side_effect=URLError("no route")):
            result = service_monitor._nginx_status()

        self.assertEqual(result["status"], "error")

    def test_get_service_status_contains_all_services(self):
        with patch("order.service_monitor._database_status"), patch(
            "order.service_monitor._redis_status"
        ), patch("order.service_monitor._nginx_status"):
            data = service_monitor.get_service_status()

        self.assertEqual({"database", "redis", "nginx"}, set(data))


class MonitorPageTests(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_superuser(
            username="monitor_admin",
            email="monitor@example.com",
            password="test-pass-123",
        )
        self.client.force_login(self.admin_user)
        self.services = {
            "database": {
                "name": "数据库",
                "status": "ok",
                "label": "正常",
                "detail": "SQLite 连接正常",
                "latency_ms": 1.2,
                "checked_at": "2026-08-01T10:00:00+08:00",
            },
            "redis": {
                "name": "Redis",
                "status": "not_configured",
                "label": "未配置",
                "detail": "未配置 REDIS_URL",
                "latency_ms": None,
                "checked_at": "2026-08-01T10:00:00+08:00",
            },
            "nginx": {
                "name": "Nginx",
                "status": "error",
                "label": "异常",
                "detail": "连接失败：timeout",
                "latency_ms": 2000.0,
                "checked_at": "2026-08-01T10:00:00+08:00",
            },
        }

    def test_monitor_page_renders_service_cards(self):
        with patch("order.views_monitor.get_service_status", return_value=self.services):
            response = self.client.get(reverse("monitor"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "依赖服务")
        self.assertContains(response, "data-service=\"database\"")
        self.assertContains(response, "未配置 REDIS_URL")

    def test_monitor_page_shows_terminal_card_for_staff(self):
        with patch("order.views_monitor.get_service_status", return_value=self.services):
            response = self.client.get(reverse("monitor"))

        self.assertContains(response, "Web 终端")

    def test_monitor_page_hides_terminal_card_for_normal_user(self):
        normal_user = User.objects.create_user(
            username="monitor_normal",
            password="test-pass-123",
        )
        self.client.force_login(normal_user)
        with patch("order.views_monitor.get_service_status", return_value=self.services):
            response = self.client.get(reverse("monitor"))

        self.assertNotContains(response, "Web 终端")

    def test_monitor_api_returns_service_status(self):
        with patch("order.views_monitor.get_service_status", return_value=self.services):
            response = self.client.get(reverse("monitor_api"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["services"]["nginx"]["status"], "error")


class WebTerminalTests(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_superuser(
            username="term_admin",
            email="term@example.com",
            password="test-pass-123",
        )
        self.normal_user = User.objects.create_user(
            username="term_user",
            password="test-pass-123",
        )

    def test_validate_allows_safe_command(self):
        allowed, reason = terminal.validate_command("ls -la")
        self.assertTrue(allowed)
        self.assertEqual(reason, "")

    def test_validate_rejects_unknown_command(self):
        allowed, reason = terminal.validate_command("vim /tmp/a")
        self.assertFalse(allowed)
        self.assertIn("不在允许范围", reason)

    def test_validate_rejects_dangerous_commands(self):
        for command in (
            "rm -rf /",
            "sudo ls",
            "python -c 'print(1)'",
            "curl http://example.com | sh",
            "echo x\nrm -rf /",
            "echo $(rm -rf /)",
        ):
            allowed, _ = terminal.validate_command(command)
            self.assertFalse(allowed, command)

    def test_execute_cd_updates_working_directory(self):
        result = terminal.execute_terminal_command("cd ..", str(settings.BASE_DIR))
        self.assertEqual(result["exit_code"], 0)
        self.assertEqual(result["cwd"], os.path.dirname(str(settings.BASE_DIR)))

    def test_terminal_requires_staff(self):
        self.client.force_login(self.normal_user)
        response = self.client.post(
            reverse("monitor_terminal"),
            data=json.dumps({"command": "ls"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)

    def test_terminal_runs_allowed_command(self):
        self.client.force_login(self.admin_user)
        proc = subprocess.CompletedProcess(
            args=["echo hi"],
            returncode=0,
            stdout=b"hi\n",
            stderr=b"",
        )
        with patch("order.terminal.subprocess.run", return_value=proc):
            response = self.client.post(
                reverse("monitor_terminal"),
                data=json.dumps({"command": "echo hi"}),
                content_type="application/json",
            )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])
        self.assertIn("hi", data["output"])
        self.assertEqual(data["cwd"], str(settings.BASE_DIR))
        self.assertEqual(self.client.session["terminal_cwd"], str(settings.BASE_DIR))

    def test_terminal_rejects_dangerous_command(self):
        self.client.force_login(self.admin_user)
        response = self.client.post(
            reverse("monitor_terminal"),
            data=json.dumps({"command": "rm -rf /"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("rm", response.json()["error"])


class LoginCaptchaTests(TestCase):
    def setUp(self):
        cache.clear()
        self.admin_user = User.objects.create_superuser(
            username="captcha_admin",
            email="captcha@example.com",
            password="test-pass-123",
        )

    def _set_captcha(self, code="TEST", expires_in=300):
        session = self.client.session
        session[CAPTCHA_SESSION_KEY] = {
            "code": code,
            "expires": time() + expires_in,
        }
        session.save()

    def test_home_page_shows_captcha(self):
        response = self.client.get(reverse("home"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "验证码")
        self.assertContains(response, 'name="captcha"')

    def test_captcha_image_sets_session_and_returns_svg(self):
        response = self.client.get(reverse("captcha_image"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("image/svg+xml", response["Content-Type"])
        self.assertIn(CAPTCHA_SESSION_KEY, self.client.session)
        payload = self.client.session[CAPTCHA_SESSION_KEY]
        self.assertEqual(len(payload["code"]), 4)
        self.assertGreater(payload["expires"], time())

    def test_login_rejects_wrong_captcha(self):
        self._set_captcha()
        response = self.client.post(
            reverse("home"),
            {
                "username": self.admin_user.username,
                "password": "test-pass-123",
                "captcha": "WRONG",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("home"))
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_login_succeeds_with_correct_captcha(self):
        self._set_captcha()
        response = self.client.post(
            reverse("home"),
            {
                "username": self.admin_user.username,
                "password": "test-pass-123",
                "captcha": "test",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("dashboard"))
        self.assertIn("_auth_user_id", self.client.session)
        self.assertNotIn(CAPTCHA_SESSION_KEY, self.client.session)

    def test_login_without_captcha_fails(self):
        response = self.client.post(
            reverse("home"),
            {
                "username": self.admin_user.username,
                "password": "test-pass-123",
                "captcha": "",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_expired_captcha_rejected(self):
        self._set_captcha(expires_in=-60)
        response = self.client.post(
            reverse("home"),
            {
                "username": self.admin_user.username,
                "password": "test-pass-123",
                "captcha": "TEST",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("home"))
        self.assertNotIn("_auth_user_id", self.client.session)
        self.assertNotIn(CAPTCHA_SESSION_KEY, self.client.session)


class ShipmentProgressTests(TestCase):
    """供应商出货后，订单进度自动调整为「剪版寄出」"""

    def setUp(self):
        self.supplier_user = User.objects.create_user(
            username="supplier_ship",
            email="supplier_ship@example.com",
            password="test-pass-123",
        )
        self.supplier = Supplier.objects.create(
            user=self.supplier_user,
            company_name="出货测试供应商",
        )
        self.admin_user = User.objects.create_superuser(
            username="admin_ship",
            email="admin_ship@example.com",
            password="test-pass-123",
        )

    def _create_order(self, order_type="sample", serial_number=1, progress_current=1):
        if order_type == "sample":
            stages = '["客户下单","通知供应商","剪版寄出","待收款","已完成"]'
        else:
            stages = '["客户下单","通知供应商","大货样","批色","查布","发货","待收款","已完成"]'
        return ClothOrder.objects.create(
            order_type=order_type,
            order_status="active",
            serial_number=serial_number,
            customer="测试客户",
            quantity_unit="码",
            price_unit="元/码",
            progress_stages=stages,
            progress_current=progress_current,
            supplier=self.supplier,
        )

    def test_model_helper_sets_sample_stage(self):
        order = self._create_order()
        self.assertTrue(order.set_progress_stage("剪版寄出"))
        order.refresh_from_db()
        self.assertEqual(order.progress_current, 2)
        self.assertEqual(order.current_stage_name, "剪版寄出")

    def test_model_helper_ignores_missing_stage(self):
        order = self._create_order(order_type="bulk", serial_number=2)
        self.assertFalse(order.set_progress_stage("剪版寄出"))
        order.refresh_from_db()
        self.assertEqual(order.progress_current, 1)

    def test_supplier_submit_price_redirects_no_shipment(self):
        order = self._create_order()
        self.client.force_login(self.supplier_user)
        response = self.client.post(
            reverse("supplier_order_detail", kwargs={"pk": order.pk}),
            {
                "finished_product_cost_price": "12.50",
                "cost_price_unit": "元/码",
                "address": "",
                "remark": "",
            },
        )
        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        self.assertFalse(order.supplier_shipped)
        self.assertEqual(str(order.finished_product_cost_price), "12.50")
        self.assertFalse(order.shipments.exists())

    def test_supplier_ajax_submit_price_only(self):
        order = self._create_order()
        self.client.force_login(self.supplier_user)
        response = self.client.post(
            reverse("supplier_order_detail", kwargs={"pk": order.pk}),
            {
                "finished_product_cost_price": "12.50",
                "cost_price_unit": "元/码",
                "address": "广州市海珠区测试路 1 号",
                "remark": "请尽快安排剪版",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "ok")
        self.assertEqual(data["cost_price"], "12.50")
        self.assertEqual(data["address"], "广州市海珠区测试路 1 号")
        self.assertEqual(data["remark"], "请尽快安排剪版")
        order.refresh_from_db()
        self.assertFalse(order.supplier_shipped)
        self.assertEqual(str(order.finished_product_cost_price), "12.50")
        self.assertEqual(order.address, "广州市海珠区测试路 1 号")
        self.assertEqual(order.remark, "请尽快安排剪版")
        self.assertFalse(order.shipments.exists())

    def test_shipment_action_buttons_use_standard_table_button_style(self):
        """出货记录操作列按钮与其他表格一致：btn btn-sm btn-outline-*，不要 py-0 px-1 的压缩样式"""
        order = self._create_order()
        order.finished_product_cost_price = "12.50"
        order.save(update_fields=["finished_product_cost_price"])
        Shipment.objects.create(
            order=order, batch_number=1, quantity="30", date="2026-08-28"
        )

        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        html = response.content.decode("utf-8")
        self.assertIn(
            'class="btn btn-sm btn-outline-primary" onclick="editShipment(', html
        )
        self.assertIn('class="btn btn-sm btn-outline-danger" title="删除"', html)
        self.assertNotIn("py-0 px-1", html)

    def test_shipment_create_blocked_without_cost_price(self):
        order = self._create_order()
        self.client.force_login(self.supplier_user)
        response = self.client.post(
            reverse("order_shipment_create", kwargs={"pk": order.pk}),
            {"date": "2026-08-28", "quantity": "30"},
        )
        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        self.assertFalse(order.supplier_shipped)
        self.assertFalse(order.shipments.exists())

    def test_supplier_price_then_shipment_advances_progress(self):
        order = self._create_order()
        self.client.force_login(self.supplier_user)
        # 先提交成品成本价格
        response = self.client.post(
            reverse("supplier_order_detail", kwargs={"pk": order.pk}),
            {
                "finished_product_cost_price": "12.50",
                "cost_price_unit": "元/码",
                "address": "",
                "remark": "",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")
        # 再新增出货（批次 1）
        response = self.client.post(
            reverse("order_shipment_create", kwargs={"pk": order.pk}),
            {"date": "2026-08-28", "quantity": "50"},
        )
        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        self.assertTrue(order.supplier_shipped)
        self.assertEqual(order.current_stage_name, "剪版寄出")
        self.assertEqual(order.total_shipment_quantity, 50)
        shipment = order.shipments.get()
        self.assertEqual(shipment.batch_number, 1)
        self.assertEqual(
            float(shipment.quantity) * float(order.finished_product_cost_price),
            625.0,
        )

    def test_shipment_create_endpoint_advances_progress(self):
        order = self._create_order()
        self.client.force_login(self.admin_user)
        response = self.client.post(
            reverse("order_shipment_create", kwargs={"pk": order.pk}),
            {"date": "2026-08-28", "quantity": "30"},
        )
        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        self.assertTrue(order.supplier_shipped)
        self.assertEqual(order.current_stage_name, "剪版寄出")

    def test_bulk_shipment_advances_to_shipped(self):
        order = self._create_order(order_type="bulk", serial_number=2)
        self.client.force_login(self.admin_user)
        response = self.client.post(
            reverse("order_shipment_create", kwargs={"pk": order.pk}),
            {"date": "2026-08-28", "quantity": "30"},
        )
        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        self.assertTrue(order.supplier_shipped)
        self.assertEqual(order.current_stage_name, "发货")

    def test_supplier_order_detail_shows_shipment_quantity_summary(self):
        order = self._create_order(serial_number=3)
        order.order_quantity = "100.00"
        order.shipment_quantity_unit = "公斤"
        order.save(update_fields=["order_quantity", "shipment_quantity_unit"])
        Shipment.objects.create(order=order, batch_number=1, date="2026-08-28", quantity="30")
        Shipment.objects.create(order=order, batch_number=2, date="2026-08-29", quantity="20")
        Shipment.objects.create(
            order=order,
            batch_number=3,
            date="2026-08-30",
            quantity="10",
            is_deleted=True,
        )

        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "出货总数量：30.00 + 20.00 = 50.00 公斤（订单数量：100.00 码）")
        self.assertContains(response, "50.00 公斤")
        self.assertNotContains(response, "50.00 米")


class SupplierExportTests(TestCase):
    def setUp(self):
        self.supplier_user = User.objects.create_user(
            username="supplier_export",
            password="test-pass-123",
        )
        self.supplier = Supplier.objects.create(
            user=self.supplier_user,
            company_name="导出测试供应商",
        )
        other_user = User.objects.create_user(
            username="other_supplier_export",
            password="test-pass-123",
        )
        self.other_supplier = Supplier.objects.create(
            user=other_user,
            company_name="其他供应商",
        )

        ClothOrder.objects.create(
            order_type="bulk",
            serial_number=101,
            order_date="2026-08-01",
            customer="导出客户A",
            cloth_type="CL-001",
            supplier=self.supplier,
        )
        ClothOrder.objects.create(
            order_type="sample",
            serial_number=102,
            order_date="2026-08-02",
            customer="导出客户B",
            cloth_type="CL-002",
            supplier=self.supplier,
        )
        ClothOrder.objects.create(
            order_type="bulk",
            serial_number=103,
            order_date="2026-08-03",
            customer="其他供应商订单",
            cloth_type="CL-003",
            supplier=self.other_supplier,
        )

    def test_supplier_can_export_current_filtered_orders_only(self):
        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_orders_export"),
            {"search": "导出客户A"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("supplier_orders_", response["Content-Disposition"])

        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook["供应商订单"].values)
        headers = list(rows[0])
        data_rows = rows[1:]
        self.assertIn("客户", headers)
        self.assertNotIn("客户付款", headers)
        self.assertEqual(len(data_rows), 1)
        self.assertEqual(data_rows[0][headers.index("客户")], "导出客户A")

    def test_supplier_export_includes_order_follower(self):
        order = ClothOrder.objects.get(serial_number=101)
        order.order_follower = "跟单小李"
        order.save(update_fields=["order_follower"])

        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_orders_export"),
            {"search": "导出客户A"},
        )

        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook["供应商订单"].values)
        headers = list(rows[0])
        self.assertIn("跟单员", headers)
        self.assertEqual(rows[1][headers.index("跟单员")], "跟单小李")

    def test_non_supplier_cannot_export_supplier_orders(self):
        admin = User.objects.create_user(
            username="supplier_export_admin",
            password="test-pass-123",
            is_staff=True,
        )
        self.client.force_login(admin)

        response = self.client.get(reverse("supplier_orders_export"))

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("dashboard"))

    def test_supplier_export_applies_shipped_status_filter(self):
        shipped_order = ClothOrder.objects.get(serial_number=102)
        shipped_order.finished_product_cost_price = "12.50"
        shipped_order.supplier_shipped = True
        shipped_order.save(update_fields=["finished_product_cost_price", "supplier_shipped"])

        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_orders_export"),
            {"status": "shipped"},
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook["供应商订单"].values)
        headers = list(rows[0])
        data_rows = rows[1:]
        self.assertEqual(len(data_rows), 1)
        self.assertEqual(data_rows[0][headers.index("序号")], 102)

    def test_supplier_dashboard_orders_by_serial_number_descending(self):
        ClothOrder.objects.create(
            order_type="bulk",
            serial_number=104,
            order_date="2026-07-01",
            customer="较早日期但较大序号",
            supplier=self.supplier,
        )
        ClothOrder.objects.create(
            order_type="bulk",
            serial_number=105,
            order_date="2026-06-01",
            customer="最早日期但最大序号",
            supplier=self.supplier,
        )

        self.client.force_login(self.supplier_user)
        response = self.client.get(reverse("supplier_dashboard"))

        self.assertEqual(response.status_code, 200)
        serials = [order.serial_number for order in response.context["orders"]]
        self.assertEqual(serials[:4], [105, 104, 102, 101])

    def test_supplier_order_detail_shows_width_and_weight_units(self):
        order = ClothOrder.objects.get(serial_number=101)
        order.width = "180.00"
        order.weight = "220.00"
        order.save(update_fields=["width", "weight"])

        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "180.00 cm")
        self.assertContains(response, "220.00 gsm")

    def test_supplier_dashboard_shows_additional_order_columns(self):
        order = ClothOrder.objects.get(serial_number=101)
        order.booth = "A01"
        order.supplier_code = "SUP-001"
        order.color = "藏青"
        order.order_quantity = "500.00"
        order.quantity_unit = "米"
        order.save(update_fields=[
            "booth",
            "supplier_code",
            "color",
            "order_quantity",
            "quantity_unit",
        ])

        self.client.force_login(self.supplier_user)
        response = self.client.get(reverse("supplier_dashboard"))

        self.assertEqual(response.status_code, 200)
        for header in ("档口", "档口编号", "颜色", "订单数量"):
            self.assertContains(response, f"<th>{header}</th>", html=True)
        self.assertContains(response, "A01")
        self.assertContains(response, "SUP-001")
        self.assertContains(response, "藏青")
        self.assertContains(response, "500.00 米")


class DatabaseBackupExcelTests(TestCase):
    def setUp(self):
        ClothOrder.objects.create(
            order_type="bulk",
            serial_number=201,
            order_date="2026-08-01",
            customer="备份客户",
            cloth_type="BACKUP-001",
            supplier=None,
        )

    def test_backup_zip_contains_database_and_both_order_excel_files(self):
        from order import db_backup

        def fake_dumpdata(*args, **kwargs):
            kwargs["stdout"].write("[]")
            return MagicMock(returncode=0)

        with tempfile.TemporaryDirectory() as temp_dir:
            with patch.object(db_backup, "BACKUP_DIR", Path(temp_dir)):
                with patch.object(db_backup.subprocess, "run", side_effect=fake_dumpdata):
                    path, method = db_backup.create_backup()

            self.assertEqual(method, "dumpdata JSON + 订单 Excel 压缩包")
            # 压缩包内同时含 JSON 与两份 Excel，故用 .zip 命名、不带 .json 后缀
            self.assertRegex(path.name, r"^db_backup_\d{8}_\d{6}\.zip$")
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                timestamp = path.name.removeprefix("db_backup_").removesuffix(".zip")
                expected_excel_sheets = {
                    f"orders_backup_{timestamp}.xlsx": "订单数据",
                    f"supplier_orders_backup_{timestamp}.xlsx": "供应商订单",
                }
                self.assertEqual(
                    names,
                    {
                        path.name.removesuffix(".zip"),
                        *expected_excel_sheets,
                    },
                )
                for excel_name, sheet_name in expected_excel_sheets.items():
                    workbook = load_workbook(BytesIO(archive.read(excel_name)), read_only=True)
                    self.assertIn(sheet_name, workbook.sheetnames)


class NotificationAccessTests(TestCase):
    """供应商消息提醒权限拦截：不能经由通知进入管理端页面"""

    def setUp(self):
        self.supplier_user = User.objects.create_user(
            username="notif_supplier",
            email="notif_supplier@example.com",
            password="test-pass-123",
        )
        self.supplier = Supplier.objects.create(
            user=self.supplier_user,
            company_name="提醒测试供应商",
        )
        self.staff_user = User.objects.create_user(
            username="notif_staff",
            email="notif_staff@example.com",
            password="test-pass-123",
            is_staff=True,
        )
        self.order = ClothOrder.objects.create(
            order_type="sample",
            order_status="active",
            serial_number=10,
            customer="测试客户",
            quantity_unit="码",
            price_unit="元/码",
            progress_stages='["客户下单","通知供应商","剪版寄出","待收款","已完成"]',
            progress_current=1,
            supplier=self.supplier,
        )

    def test_supplier_blocked_from_admin_dashboard(self):
        self.client.force_login(self.supplier_user)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("supplier_dashboard"))

    def test_supplier_blocked_from_admin_order_pages(self):
        self.client.force_login(self.supplier_user)
        response = self.client.get(reverse("order_list"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("supplier_dashboard"))
        response = self.client.get(reverse("order_detail", kwargs={"pk": self.order.pk}))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("supplier_dashboard"))

    def test_supplier_notification_list_uses_supplier_base(self):
        Notification.objects.create(
            recipient=self.supplier_user,
            title="测试提醒",
            message="订单测试",
            link=f"/orders/{self.order.pk}/",
        )
        self.client.force_login(self.supplier_user)
        response = self.client.get(reverse("notification_list"), HTTP_HOST="localhost")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertIn("供应商面板", html)
        self.assertNotIn("订单管理", html)
        self.assertIn(f"/supplier/orders/{self.order.pk}/", html)

    def test_admin_notification_list_uses_admin_base(self):
        Notification.objects.create(
            recipient=self.staff_user,
            title="测试提醒",
            link=f"/orders/{self.order.pk}/",
        )
        self.client.force_login(self.staff_user)
        response = self.client.get(reverse("notification_list"), HTTP_HOST="localhost")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertIn("订单管理", html)
        self.assertIn(f"/orders/{self.order.pk}/", html)
        self.assertNotIn("供应商面板", html)

    def test_supplier_unread_list_link_rewritten(self):
        Notification.objects.create(
            recipient=self.supplier_user,
            title="测试提醒",
            link=f"/orders/{self.order.pk}/",
        )
        self.client.force_login(self.supplier_user)
        response = self.client.get(reverse("notification_unread_list"))
        data = response.json()
        self.assertEqual(
            data["notifications"][0]["link"],
            f"/supplier/orders/{self.order.pk}/",
        )

    def test_admin_notification_link_unchanged(self):
        Notification.objects.create(
            recipient=self.staff_user,
            title="测试提醒",
            link=f"/orders/{self.order.pk}/",
        )
        self.client.force_login(self.staff_user)
        response = self.client.get(reverse("notification_unread_list"))
        data = response.json()
        self.assertEqual(data["notifications"][0]["link"], f"/orders/{self.order.pk}/")

    def test_shipment_create_notifies_staff(self):
        self.order.finished_product_cost_price = "12.50"
        self.order.save(update_fields=["finished_product_cost_price"])
        self.client.force_login(self.supplier_user)
        response = self.client.post(
            reverse("order_shipment_create", kwargs={"pk": self.order.pk}),
            {"date": "2026-08-28", "quantity": "50"},
        )
        self.assertEqual(response.status_code, 302)
        notif = self.staff_user.notifications.first()
        self.assertIsNotNone(notif)
        self.assertIn("已对订单 #10 提交出货", notif.message)
        self.assertEqual(notif.link, f"/orders/{self.order.pk}/")

    def test_mark_all_read_redirects_instead_of_json(self):
        for i in range(2):
            Notification.objects.create(
                recipient=self.staff_user,
                title=f"测试{i}",
                link=f"/orders/{self.order.pk}/",
            )
        self.client.force_login(self.staff_user)
        list_url = reverse("notification_list")
        response = self.client.post(
            reverse("notification_mark_all_read"),
            HTTP_REFERER=list_url,
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, list_url)
        self.assertEqual(self.staff_user.notifications.filter(is_read=False).count(), 0)

    def test_mark_read_form_redirects_ajax_returns_json(self):
        n = Notification.objects.create(
            recipient=self.staff_user,
            title="测试提醒",
            link=f"/orders/{self.order.pk}/",
        )
        self.client.force_login(self.staff_user)
        # 表单提交：刷新回来源页
        response = self.client.post(
            reverse("notification_mark_read", kwargs={"pk": n.pk}),
            HTTP_REFERER=reverse("notification_list"),
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("notification_list"))
        n.refresh_from_db()
        self.assertTrue(n.is_read)
        # AJAX：返回 JSON
        n2 = Notification.objects.create(
            recipient=self.staff_user,
            title="测试提醒2",
        )
        response = self.client.post(
            reverse("notification_mark_read", kwargs={"pk": n2.pk}),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"ok": True})
        n2.refresh_from_db()
        self.assertTrue(n2.is_read)


class ErrorPageTests(TestCase):
    """统一错误/异常提示页"""

    def setUp(self):
        self.admin_user = User.objects.create_superuser(
            username="error_admin",
            email="error_admin@example.com",
            password="test-pass-123",
        )
        self.supplier_user = User.objects.create_user(
            username="error_supplier",
            email="error_supplier@example.com",
            password="test-pass-123",
        )
        self.supplier = Supplier.objects.create(
            user=self.supplier_user,
            company_name="错误页测试供应商",
        )
        self.factory = RequestFactory()

    def _request(self, user=None, ajax=False):
        kwargs = {"HTTP_X_REQUESTED_WITH": "XMLHttpRequest"} if ajax else {}
        req = self.factory.get("/some/error/path/", **kwargs)
        req.user = user if user is not None else AnonymousUser()
        return req

    def test_handlers_return_expected_status(self):
        self.assertEqual(handler400(self._request()).status_code, 400)
        self.assertEqual(handler403(self._request()).status_code, 403)
        self.assertEqual(handler404(self._request()).status_code, 404)
        self.assertEqual(handler500(self._request()).status_code, 500)

    def test_supplier_error_page_uses_supplier_base(self):
        response = handler404(self._request(user=self.supplier_user))
        self.assertEqual(response.status_code, 404)
        html = response.content.decode("utf-8")
        self.assertIn("供应商面板", html)
        self.assertIn("404", html)
        self.assertIn("您访问的页面不存在或已被移除。", html)
        self.assertNotIn("订单管理", html)

    def test_admin_error_page_uses_admin_base(self):
        response = handler404(self._request(user=self.admin_user))
        self.assertEqual(response.status_code, 404)
        html = response.content.decode("utf-8")
        self.assertIn("订单管理", html)
        self.assertNotIn("供应商面板", html)

    def test_anonymous_error_page_renders(self):
        response = handler403(self._request())
        self.assertEqual(response.status_code, 403)
        html = response.content.decode("utf-8")
        self.assertIn("403", html)

    def test_ajax_error_returns_json(self):
        response = handler404(self._request(user=self.admin_user, ajax=True))
        self.assertEqual(response.status_code, 404)
        data = json.loads(response.content)
        self.assertEqual(data["code"], "404")
        self.assertIn("页面不存在", data["error"])

    @override_settings(DEBUG=False, ALLOWED_HOSTS=["testserver"])
    def test_urlconf_404_wiring(self):
        self.client.force_login(self.admin_user)
        response = self.client.get("/this-page-does-not-exist/")
        self.assertEqual(response.status_code, 404)
        self.assertContains(response, "404", status_code=404)
        self.assertContains(response, "您访问的页面不存在或已被移除。", status_code=404)


# 管理端与供应商端共用同一份导航栏样式，两端渲染出来的这段 CSS 必须完全一致
SHARED_NAVBAR_CSS = """.navbar {
    background: var(--gradient-primary) !important;
    box-shadow: 0 4px 20px rgba(67, 97, 238, 0.3);
    padding: 1rem 0;
}"""


class OrderTypeBackgroundTests(TestCase):
    """订单详情/编辑/创建页背景色随订单类型变化"""

    def setUp(self):
        self.admin_user = User.objects.create_superuser(
            username="order_bg_admin",
            email="order_bg_admin@example.com",
            password="test-pass-123",
        )
        self.client.force_login(self.admin_user)
        self._serial = 700

    def _create_order(self, order_type):
        self._serial += 1
        return ClothOrder.objects.create(
            order_type=order_type,
            order_status="active",
            serial_number=self._serial,
            customer="背景色客户",
            cloth_type="BG-001",
        )

    def test_detail_page_body_class_matches_order_type(self):
        for order_type in ["bulk", "bulk_print", "sample", "other"]:
            order = self._create_order(order_type)
            response = self.client.get(
                reverse("order_detail", kwargs={"pk": order.pk})
            )
            self.assertContains(response, f'<body class="order-type-{order_type}">')

    def test_detail_edit_mode_keeps_order_type_background(self):
        order = self._create_order("sample")
        response = self.client.get(
            reverse("order_detail", kwargs={"pk": order.pk}),
            {"edit": "1"},
        )
        self.assertContains(response, '<body class="order-type-sample">')
        # 编辑态必须能改订单类型，否则背景色无从切换
        self.assertContains(response, 'id="id_order_type"')

    def test_create_page_defaults_to_bulk_background(self):
        response = self.client.get(reverse("order_create"))
        self.assertContains(response, '<body class="order-type-bulk">')

    def test_detail_has_grouped_order_nav_card(self):
        bulk_order = self._create_order("bulk")
        sample_order = self._create_order("sample")

        response = self.client.get(
            reverse("order_detail", kwargs={"pk": bulk_order.pk})
        )

        html = response.content.decode("utf-8")
        self.assertIn("订单导航", html)
        # 分类是折叠式下拉，触发器指向各自的折叠体
        self.assertIn('href="#navGroup-bulk"', html)
        self.assertIn('href="#navGroup-sample"', html)
        # 当前订单置顶高亮，且不出现在分类列表里（不重复）
        self.assertEqual(
            html.count(
                'class="list-group-item list-group-item-action py-2 sidebar-order-active"'
            ),
            1,
        )
        self.assertEqual(html.count('<span class="badge bg-primary">当前</span>'), 1)
        self.assertEqual(
            html.count(f'href="{reverse("order_detail", kwargs={"pk": bulk_order.pk})}"'),
            1,
        )
        # 其他分类的订单可以点进去
        self.assertIn(
            f'href="{reverse("order_detail", kwargs={"pk": sample_order.pk})}"', html
        )

    def test_order_nav_categories_collapsed_by_default(self):
        bulk_order = self._create_order("bulk")
        self._create_order("sample")

        response = self.client.get(
            reverse("order_detail", kwargs={"pk": bulk_order.pk})
        )

        html = response.content.decode("utf-8")
        # 触发器可折叠，且默认收起（折叠体是 collapse，不带 show）
        self.assertIn(
            'data-bs-toggle="collapse" href="#navGroup-bulk" role="button" aria-expanded="false"',
            html,
        )
        self.assertIn('<div class="collapse" id="navGroup-bulk">', html)
        self.assertNotIn('collapse show" id="navGroup-bulk"', html)

    def test_order_nav_rows_show_customer_follower_supplier(self):
        order = self._create_order("bulk")
        ClothOrder.objects.filter(pk=order.pk).update(
            order_follower="跟单老陈",
            finished_product_supplier="成品供应商A",
        )

        response = self.client.get(reverse("order_detail", kwargs={"pk": order.pk}))

        html = response.content.decode("utf-8")
        self.assertIn("sidebar-order-meta", html)
        self.assertIn(
            "客户：背景色客户｜跟单员：跟单老陈｜供应商：成品供应商A", html
        )

    def test_order_nav_groups_skip_empty_types(self):
        order = self._create_order("bulk")

        response = self.client.get(reverse("order_detail", kwargs={"pk": order.pk}))

        html = response.content.decode("utf-8")
        self.assertIn('href="#navGroup-bulk"', html)
        # 没有订单的分类不渲染（导航栏下拉菜单里的分类链接不算）
        self.assertNotIn('href="#navGroup-bulk_print"', html)
        self.assertNotIn('href="#navGroup-sample"', html)

    def test_order_nav_group_caps_at_20_recent(self):
        orders = [self._create_order("bulk") for _ in range(25)]
        current = self._create_order("bulk")

        response = self.client.get(reverse("order_detail", kwargs={"pk": current.pk}))

        html = response.content.decode("utf-8")
        newest = orders[-1]
        excluded = orders[-21]
        self.assertIn(
            f'href="{reverse("order_detail", kwargs={"pk": newest.pk})}"', html
        )
        self.assertNotIn(
            f'href="{reverse("order_detail", kwargs={"pk": excluded.pk})}"', html
        )

    def test_admin_page_includes_shared_navbar_theme(self):
        response = self.client.get(reverse("order_list"))
        html = response.content.decode("utf-8")
        self.assertIn(SHARED_NAVBAR_CSS, html)
        self.assertIn(
            "--gradient-primary: linear-gradient(135deg, #4361ee 0%, #3a0ca3 100%)",
            html,
        )

    def test_create_page_order_type_select_shows_badge(self):
        response = self.client.get(reverse("order_create"))
        html = response.content.decode("utf-8")
        self.assertEqual(html.count('id="orderTypeBadge"'), 1)
        # 徽章紧跟在订单类型下拉框后，初始值与表单默认值一致
        self.assertRegex(html, r'id="orderTypeBadge"[^>]*>大货订单<')
        self.assertIn("order-type-badge-bulk", html)

    def test_edit_page_order_type_select_shows_badge(self):
        order = self._create_order("bulk_print")
        response = self.client.get(
            reverse("order_detail", kwargs={"pk": order.pk}),
            {"edit": "1"},
        )
        html = response.content.decode("utf-8")
        self.assertEqual(html.count('id="orderTypeBadge"'), 1)
        self.assertRegex(html, r'id="orderTypeBadge"[^>]*>大货-印花<')
        self.assertIn("order-type-badge-bulk_print", html)

    def test_badge_hidden_when_no_order_type_selected(self):
        """未选类型时不渲染空徽章"""
        from django.template.loader import render_to_string

        from order.forms import ClothOrderForm

        form = ClothOrderForm(initial={"order_type": ""})
        html = render_to_string(
            "order/_order_type_field.html", {"field": form["order_type"]}
        )
        self.assertIn('id="orderTypeBadge"', html)
        self.assertIn("d-none", html)

    def test_badge_and_background_share_the_same_colour(self):
        """徽章与页面底色必须引用同一个颜色变量，避免两处配色走偏"""
        response = self.client.get(reverse("order_create"))
        html = response.content.decode("utf-8")
        for order_type in ["bulk", "bulk_print", "sample", "other"]:
            self.assertIn(f"body.order-type-{order_type} {{", html)
            self.assertIn(f".badge.order-type-badge-{order_type} {{", html)
            self.assertEqual(html.count(f"var(--order-tint-{order_type})"), 2)

    def test_detail_page_shows_coloured_type_badge(self):
        order = self._create_order("sample")
        response = self.client.get(
            reverse("order_detail", kwargs={"pk": order.pk})
        )
        self.assertContains(response, "badge order-type-badge order-type-badge-sample")

    def test_order_list_badge_uses_type_palette(self):
        self._create_order("bulk_print")
        response = self.client.get(reverse("order_list"))
        html = response.content.decode("utf-8")
        self.assertIn("badge order-type-badge order-type-badge-bulk_print", html)
        # 旧的临时配色不应再挂在订单类型上
        self.assertNotIn("order_type == 'bulk_print' %}bg-info", html)

    def test_dashboard_badge_uses_type_palette(self):
        self._create_order("sample")
        response = self.client.get(reverse("dashboard"))
        self.assertContains(response, "badge order-type-badge order-type-badge-sample")

    def test_blank_order_type_falls_back_to_other_badge(self):
        order = self._create_order("bulk")
        ClothOrder.objects.filter(pk=order.pk).update(order_type="")
        response = self.client.get(
            reverse("order_detail", kwargs={"pk": order.pk})
        )
        self.assertContains(response, "order-type-badge-other")

    def test_pages_without_order_type_keep_default_background(self):
        response = self.client.get(reverse("order_list"))
        self.assertContains(response, '<body class="">')


class SupplierOrderTypeThemeTests(TestCase):
    """供应商端：列表用订单类型徽章，详情页按订单类型换底色"""

    def setUp(self):
        self.supplier_user = User.objects.create_user(
            username="supplier_theme",
            email="supplier_theme@example.com",
            password="test-pass-123",
        )
        self.supplier = Supplier.objects.create(
            user=self.supplier_user,
            company_name="配色测试供应商",
        )
        self.client.force_login(self.supplier_user)
        self._serial = 800

    def _create_order(self, order_type):
        self._serial += 1
        return ClothOrder.objects.create(
            order_type=order_type,
            order_status="active",
            serial_number=self._serial,
            customer="配色客户",
            cloth_type="THEME-001",
            supplier=self.supplier,
        )

    def test_supplier_page_uses_same_navbar_theme_as_admin(self):
        """供应商端横幅必须与管理端同一份样式：品牌渐变、导航项内边距/悬浮位移、下拉面板阴影"""
        response = self.client.get(reverse("supplier_dashboard"))
        html = response.content.decode("utf-8")
        self.assertIn(SHARED_NAVBAR_CSS, html)
        self.assertIn("padding: 0.5rem 1rem", html)
        self.assertIn("transform: translateY(-2px)", html)
        self.assertIn("box-shadow: 0 10px 30px rgba(0, 0, 0, 0.1)", html)
        # 旧的纯色横横幅不应残留
        self.assertNotIn("#4361ee !important", html)

    def test_sidebar_categories_are_collapsed_dropdowns(self):
        order = self._create_order("sample")

        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        html = response.content.decode("utf-8")
        self.assertIn(
            'data-bs-toggle="collapse" href="#statusGroup-pending" role="button" aria-expanded="false"',
            html,
        )
        self.assertIn('<div class="collapse" id="statusGroup-pending">', html)

    def test_sidebar_rows_show_customer_and_follower(self):
        order = self._create_order("sample")
        ClothOrder.objects.filter(pk=order.pk).update(order_follower="跟单小王")

        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        html = response.content.decode("utf-8")
        self.assertIn("sidebar-order-meta", html)
        self.assertIn("客户：配色客户｜跟单员：跟单小王", html)

    def test_dashboard_order_type_column_uses_badge(self):
        self._create_order("bulk_print")
        response = self.client.get(reverse("supplier_dashboard"))
        html = response.content.decode("utf-8")
        self.assertIn("badge order-type-badge order-type-badge-bulk_print", html)
        self.assertIn("大货-印花", html)

    def test_detail_info_card_shows_type_badge(self):
        order = self._create_order("sample")
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )
        self.assertContains(response, "badge order-type-badge order-type-badge-sample")

    def test_sidebar_lists_only_this_supplier_orders(self):
        other_user = User.objects.create_user(
            username="sidebar_other", password="test-pass-123"
        )
        other_supplier = Supplier.objects.create(
            user=other_user, company_name="别家供应商"
        )
        ClothOrder.objects.create(
            order_type="bulk",
            order_status="active",
            serial_number=9999,
            customer="别家客户",
            cloth_type="OTHER-001",
            supplier=other_supplier,
        )
        order = self._create_order("sample")

        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        html = response.content.decode("utf-8")
        self.assertNotIn("#9999", html)
        self.assertIn("查看全部订单", html)

    def test_sidebar_highlights_current_order_and_lists_recent(self):
        previous = self._create_order("bulk")
        current = self._create_order("sample")

        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": current.pk})
        )

        html = response.content.decode("utf-8")
        current_row = (
            'class="list-group-item list-group-item-action py-2 sidebar-order-active"'
        )
        self.assertEqual(html.count(current_row), 1)
        self.assertIn("当前", html)
        # 当前订单在卡片顶部，分类列表在它下面
        self.assertLess(html.index(current_row), html.index('href="#statusGroup-'))
        self.assertIn(f"#{current.serial_number}", html)
        self.assertIn(f"#{previous.serial_number}", html)

    def test_sidebar_shows_type_and_state_badges(self):
        order = self._create_order("bulk_print")

        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        html = response.content.decode("utf-8")
        self.assertIn("badge order-type-badge order-type-badge-bulk_print", html)
        self.assertIn("待提交", html)

    def test_sidebar_state_badge_reflects_shipped_and_quoted(self):
        quoted = self._create_order("bulk")
        ClothOrder.objects.filter(pk=quoted.pk).update(
            finished_product_cost_price="10.00"
        )
        shipped = self._create_order("bulk")
        ClothOrder.objects.filter(pk=shipped.pk).update(
            finished_product_cost_price="10.00", supplier_shipped=True
        )
        Shipment.objects.create(order=shipped, batch_number=1, quantity="30")

        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": quoted.pk})
        )

        html = response.content.decode("utf-8")
        self.assertIn("待出货", html)
        self.assertIn("已出货", html)

    def test_detail_page_body_class_matches_order_type(self):
        for order_type in ["bulk", "bulk_print", "sample", "other"]:
            order = self._create_order(order_type)
            response = self.client.get(
                reverse("supplier_order_detail", kwargs={"pk": order.pk})
            )
            self.assertContains(response, f'<body class="order-type-{order_type}">')

    def test_supplier_base_carries_the_shared_palette(self):
        """供应商端不继承 base.html，必须显式引入共享色板，否则底色与徽章无样式"""
        response = self.client.get(reverse("supplier_dashboard"))
        html = response.content.decode("utf-8")
        for order_type in ["bulk", "bulk_print", "sample", "other"]:
            self.assertIn(f"body.order-type-{order_type} {{", html)
            self.assertIn(f".badge.order-type-badge-{order_type} {{", html)
            self.assertEqual(html.count(f"var(--order-tint-{order_type})"), 2)

    def test_dashboard_keeps_default_background(self):
        """列表面板混有多种订单类型，底色保持默认"""
        self._create_order("sample")
        response = self.client.get(reverse("supplier_dashboard"))
        self.assertContains(response, '<body class="">')


class SupplierPaginationTests(TestCase):
    """供应商端订单列表分页与管理端订单列表用同一套 Vue 组件"""

    def setUp(self):
        self.supplier_user = User.objects.create_user(
            username="supplier_pagination",
            email="supplier_pagination@example.com",
            password="test-pass-123",
        )
        self.supplier = Supplier.objects.create(
            user=self.supplier_user,
            company_name="分页测试供应商",
        )
        self.client.force_login(self.supplier_user)
        ClothOrder.objects.bulk_create([
            ClothOrder(
                order_type="bulk",
                order_status="active",
                serial_number=900 + index,
                customer=f"分页客户{index}",
                cloth_type="PAGE-001",
                supplier=self.supplier,
            )
            for index in range(25)
        ])

    def test_pagination_uses_the_shared_admin_component(self):
        response = self.client.get(reverse("supplier_dashboard"))
        html = response.content.decode("utf-8")
        self.assertIn('id="order-pagination-vue"', html)
        self.assertIn('id="order-pagination-data"', html)
        # 管理端 order_list 用的就是这个包装脚本与共享组件
        self.assertIn("app.mount('#order-pagination-vue')", html)
        self.assertIn("window.__paginationBarComponent", html)
        # 旧的服务端分页组件（含每页选择器）不应再出现在供应商端
        self.assertNotIn('<nav class="order-pagination"', html)
        self.assertNotIn('name="per_page"', html)

    def test_pagination_json_tracks_current_page(self):
        response = self.client.get(reverse("supplier_dashboard"))
        self.assertEqual(response.context["order_pagination_json"]["current_page"], 1)
        self.assertEqual(response.context["order_pagination_json"]["total_pages"], 2)

        response = self.client.get(reverse("supplier_dashboard"), {"page": 2})
        self.assertEqual(response.context["order_pagination_json"]["current_page"], 2)

    def test_out_of_range_page_returns_404_like_admin(self):
        """越界页码走 Django ListView 默认行为（404），与管理端订单列表一致；
        前端组件的「跳转」框本身有范围校验，正常点选不会越界。"""
        response = self.client.get(reverse("supplier_dashboard"), {"page": 999})
        self.assertEqual(response.status_code, 404)


class SupplierStatusClassificationTests(TestCase):
    """供应商端状态分类：待提交 / 待出货 / 已出货
    —— 分类展示在订单详情页左侧「我的订单」卡片里，面板的 status 筛选复用同一套判定"""

    def setUp(self):
        self.supplier_user = User.objects.create_user(
            username="supplier_status",
            email="supplier_status@example.com",
            password="test-pass-123",
        )
        self.supplier = Supplier.objects.create(
            user=self.supplier_user, company_name="分类测试供应商"
        )
        self.client.force_login(self.supplier_user)

        self.pending = ClothOrder.objects.create(
            order_type="bulk", order_status="active", serial_number=601,
            customer="待提交客户", cloth_type="ST-001", supplier=self.supplier,
        )
        self.quoted = ClothOrder.objects.create(
            order_type="bulk", order_status="active", serial_number=602,
            customer="待出货客户", cloth_type="ST-002", supplier=self.supplier,
            finished_product_cost_price="10.00",
        )
        self.shipped = ClothOrder.objects.create(
            order_type="bulk", order_status="active", serial_number=603,
            customer="已出货客户", cloth_type="ST-003", supplier=self.supplier,
            finished_product_cost_price="10.00", supplier_shipped=True,
        )
        Shipment.objects.create(order=self.shipped, batch_number=1, quantity="20")

    def test_detail_sidebar_shows_three_status_categories(self):
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": self.pending.pk})
        )

        html = response.content.decode("utf-8")
        for value in ("pending", "quoted", "shipped"):
            self.assertIn(
                f'data-bs-toggle="collapse" href="#statusGroup-{value}" role="button" aria-expanded="false"',
                html,
            )
            self.assertIn(f'<div class="collapse" id="statusGroup-{value}">', html)
        for label in ("待提交", "待出货", "已出货"):
            self.assertIn(label, html)

    def test_detail_sidebar_puts_each_order_in_its_status_group(self):
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": self.pending.pk})
        )

        groups = {
            group["value"]: [order.pk for order in group["orders"]]
            for group in response.context["status_groups"]
        }
        # 当前订单置顶显示，不再出现在分类列表里
        self.assertEqual(groups["pending"], [])
        self.assertEqual(groups["quoted"], [self.quoted.pk])
        self.assertEqual(groups["shipped"], [self.shipped.pk])

        totals = {
            group["value"]: group["total"]
            for group in response.context["status_groups"]
        }
        self.assertEqual(totals, {"pending": 1, "quoted": 1, "shipped": 1})

    def test_detail_sidebar_skips_empty_status_group(self):
        Shipment.objects.filter(order=self.shipped).delete()
        ClothOrder.objects.filter(pk=self.shipped.pk).update(supplier_shipped=False)

        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": self.pending.pk})
        )

        html = response.content.decode("utf-8")
        self.assertNotIn('href="#statusGroup-shipped"', html)
        self.assertIn('href="#statusGroup-quoted"', html)

    def test_detail_sidebar_only_lists_this_supplier(self):
        other_user = User.objects.create_user(
            username="status_other", password="test-pass-123"
        )
        other_supplier = Supplier.objects.create(
            user=other_user, company_name="别家供应商"
        )
        ClothOrder.objects.create(
            order_type="bulk", order_status="active", serial_number=8888,
            customer="别家客户", cloth_type="OTH-001", supplier=other_supplier,
        )

        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": self.pending.pk})
        )

        listed = [
            order.pk
            for group in response.context["status_groups"]
            for order in group["orders"]
        ]
        self.assertNotIn(8888, [order.serial_number for order in ClothOrder.objects.filter(pk__in=listed)])
        self.assertNotIn("#8888", response.content.decode("utf-8"))

    def test_panel_status_filter_matches_the_same_three_categories(self):
        """面板筛选与详情页分类共用同一套判定"""
        for status, expected in (
            ("pending", self.pending),
            ("quoted", self.quoted),
            ("shipped", self.shipped),
        ):
            response = self.client.get(
                reverse("supplier_dashboard"), {"status": status}
            )
            self.assertEqual(
                [order.pk for order in response.context["orders"]],
                [expected.pk],
                status,
            )

    def test_panel_category_counts_stay_stable_when_filtered(self):
        """筛选某一档时其他档的计数不能归零，列表头部只显示当前筛选条数"""
        response = self.client.get(reverse("supplier_dashboard"), {"status": "shipped"})

        summary = response.context["supplier_summary_json"]
        self.assertEqual(summary["total_count"], 3)
        self.assertEqual(summary["pending_count"], 1)
        self.assertEqual(summary["quoted_wait_ship_count"], 1)
        self.assertEqual(summary["shipped_count"], 1)
        self.assertEqual(response.context["total_count"], 1)


class CostPriceHistoryTests(TestCase):
    """成品成本价格：同布种最近成本价自动填写 + 每次改价留痕并展示时间"""

    def setUp(self):
        self.supplier_user = User.objects.create_user(
            username="cost_price_supplier",
            email="cost_price_supplier@example.com",
            password="test-pass-123",
        )
        self.supplier = Supplier.objects.create(
            user=self.supplier_user,
            company_name="成本价测试供应商",
        )
        self.admin_user = User.objects.create_superuser(
            username="cost_price_admin",
            email="cost_price_admin@example.com",
            password="test-pass-123",
        )

    def _create_order(self, serial_number=1, cloth_type="SYL-001", follower="老陈", supplier=None):
        return ClothOrder.objects.create(
            order_type="bulk",
            order_status="active",
            serial_number=serial_number,
            order_date="2026-08-01",
            customer="成本价客户",
            cloth_type=cloth_type,
            order_follower=follower,
            quantity_unit="米",
            price_unit="元/米",
            progress_stages='["客户下单","通知供应商","大货样","批色","查布","发货","待收款","已完成"]',
            progress_current=1,
            supplier=supplier or self.supplier,
        )

    def _set_price_without_history(self, order, price, unit="元/米"):
        """模拟功能上线前的历史数据：直接写库，不产生改价记录"""
        ClothOrder.objects.filter(pk=order.pk).update(
            finished_product_cost_price=price,
            cost_price_unit=unit,
        )

    def _submit_price(self, order, price):
        return self.client.post(
            reverse("supplier_order_detail", kwargs={"pk": order.pk}),
            {
                "finished_product_cost_price": price,
                "cost_price_unit": "元/米",
                "address": "",
                "remark": "",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

    def test_each_price_change_is_recorded_with_time(self):
        order = self._create_order()
        self.client.force_login(self.supplier_user)

        self._submit_price(order, "12.50")
        self._submit_price(order, "13.00")

        records = list(order.cost_price_history.all())
        # 默认按修改时间倒序，最新价格在最前
        self.assertEqual([str(r.price) for r in records], ["13.00", "12.50"])
        self.assertEqual([r.unit for r in records], ["元/米", "元/米"])
        self.assertEqual(records[0].changed_by, self.supplier_user)
        self.assertIsNotNone(records[0].created_at)

    def test_resubmitting_same_price_does_not_add_record(self):
        order = self._create_order()
        self.client.force_login(self.supplier_user)

        self._submit_price(order, "12.50")
        self._submit_price(order, "12.50")

        self.assertEqual(order.cost_price_history.count(), 1)

    def test_ajax_response_returns_updated_history_html(self):
        order = self._create_order()
        self.client.force_login(self.supplier_user)

        response = self._submit_price(order, "12.50")

        data = response.json()
        self.assertEqual(data["status"], "ok")
        self.assertIn("12.50", data["cost_price_history_html"])
        self.assertIn("修改时间", data["cost_price_history_html"])

    def test_same_cloth_type_price_is_prefilled(self):
        legacy = self._create_order(serial_number=2, cloth_type="SYL-777")
        self._set_price_without_history(legacy, "11.00")
        order = self._create_order(serial_number=3, cloth_type="SYL-777")

        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        self.assertEqual(response.status_code, 200)
        reference = response.context["cost_price_reference"]
        self.assertEqual(str(reference["price"]), "11.00")
        self.assertEqual(reference["unit"], "元/米")
        html = response.content.decode("utf-8")
        self.assertIn('value="11.00"', html)
        self.assertIn("同布种最近成本价", html)

    def test_reference_ignores_other_cloth_type_and_other_supplier(self):
        other_user = User.objects.create_user(
            username="cost_price_other_supplier",
            password="test-pass-123",
        )
        other_supplier = Supplier.objects.create(
            user=other_user,
            company_name="别家供应商",
        )
        same_cloth_other_supplier = self._create_order(
            serial_number=4, cloth_type="SYL-888", supplier=other_supplier
        )
        self._set_price_without_history(same_cloth_other_supplier, "9.90")

        other_cloth_same_supplier = self._create_order(
            serial_number=5, cloth_type="SYL-999"
        )
        self._set_price_without_history(other_cloth_same_supplier, "8.80")

        order = self._create_order(serial_number=6, cloth_type="SYL-888")

        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        self.assertIsNone(response.context["cost_price_reference"])

    def test_no_reference_needed_when_price_already_submitted(self):
        order = self._create_order(cloth_type="SYL-555")
        self._set_price_without_history(order, "7.70")

        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        self.assertIsNone(response.context["cost_price_reference"])

    def test_admin_save_records_price_change_operator(self):
        from django.contrib.admin.sites import AdminSite

        from order.admin import ClothOrderAdmin

        order = self._create_order()
        self._set_price_without_history(order, "10.00")
        order.refresh_from_db()

        order.finished_product_cost_price = "11.00"
        request = RequestFactory().post("/admin/order/clothorder/")
        request.user = self.admin_user
        ClothOrderAdmin(ClothOrder, AdminSite()).save_model(
            request, order, form=None, change=True
        )

        record = order.cost_price_history.get()
        self.assertEqual(str(record.price), "11.00")
        self.assertEqual(record.changed_by, self.admin_user)

    def test_history_panel_sits_under_the_cost_price_field(self):
        """改价记录紧跟在成品成本价格字段下面，不再留在左侧栏"""
        order = self._create_order()
        self.client.force_login(self.supplier_user)

        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        html = response.content.decode("utf-8")
        price_pos = html.index('id="id_finished_product_cost_price"')
        history_pos = html.index('id="costPriceHistoryWrap"')
        submit_pos = html.index('id="priceSubmitBtn"')
        self.assertLess(price_pos, history_pos)
        self.assertLess(history_pos, submit_pos)
        self.assertEqual(html.count('id="costPriceHistoryWrap"'), 1)

    def test_admin_order_detail_shows_price_history(self):
        order = self._create_order()
        self.client.force_login(self.supplier_user)
        self._submit_price(order, "12.50")

        self.client.force_login(self.admin_user)
        response = self.client.get(reverse("order_detail", kwargs={"pk": order.pk}))

        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")
        self.assertIn("成品成本价格修改记录", html)
        self.assertIn("12.50", html)
        self.assertIn("cost_price_supplier", html)

    def test_shipment_form_defaults_to_today(self):
        self.assertEqual(ShipmentForm().fields["date"].initial, timezone.localdate())

    def test_supplier_detail_shipment_date_defaults_to_today(self):
        order = self._create_order()

        self.client.force_login(self.supplier_user)
        response = self.client.get(
            reverse("supplier_order_detail", kwargs={"pk": order.pk})
        )

        html = response.content.decode("utf-8")
        expected = f'id="id_date" value="{timezone.localdate().isoformat()}"'
        self.assertIn(expected, html)

    def test_supplier_dashboard_lists_order_follower(self):
        self._create_order(follower="跟单小王")

        self.client.force_login(self.supplier_user)
        response = self.client.get(reverse("supplier_dashboard"))

        html = response.content.decode("utf-8")
        self.assertIn("<th>跟单员</th>", html)
        self.assertIn("跟单小王", html)
